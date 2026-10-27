from conftest import without_setup
from conftest import comparison_wire
from conftest import onboard_test_account
import io
import json
import time
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader
from docx import Document
from app import applications, comparison, db, documents, config, ai
from app.main import app
from app.network import ServiceError
from test_app import fixture_analysis


def wait_task(client,response):
    assert response.status_code==200,response.text
    # Real PDF fitting can take longer on a shared ARM runner. Bound elapsed
    # time instead of assuming a fixed number of fast HTTP polls is sufficient.
    task_id=response.json()['task_id'];deadline=time.monotonic()+120;task=None
    while time.monotonic()<deadline:
        task=client.get('/api/tasks/'+task_id).json()
        if task['status'] not in ('queued','running'): return task
        time.sleep(.1)
    raise AssertionError(f'Task did not finish within 120 seconds: {task}')


def req(status='supported',importance='required',kind='qualification'):
    return {'keyword':'Python','status':status,'importance':importance,'kind':kind}

@pytest.mark.parametrize('rows,expected',[
    ([req()],100),([req('partial')],50),([req('unknown')],0),
    ([req(),req('unknown','preferred')],70),([req('partial'),req('supported','preferred')],65),
    ([req('supported','preferred')],100),([req(kind='eligibility')],None),([],None),([req(importance='unclear')],None)
])
def test_qualification_scoring(rows,expected):
    assert comparison.score(rows)['value']==expected


def test_private_authorization_backup_and_resume_exclusion(client,sample,account):
    private={'authorized_us':'yes','sponsorship_now':'no','sponsorship_future':'unsure'}
    assert client.put('/api/settings',json=private).status_code==200
    ids=[r['id'] for r in without_setup(db.list_entities(account['id'],'experience'))]
    response=client.post('/api/resumes',json={'experience_ids':ids,'job_id':sample['id']})
    assert response.status_code==200,response.text
    version=response.json()
    for k in private: assert k not in json.dumps(version['data'])
    assert 'artifacts' not in version['data']
    assert 'artifacts' not in json.dumps(client.get('/api/workspace').json())
    backup=client.get('/api/account/export').json()
    assert all(k not in backup['profile'] for k in private)
    assert all(db.settings(account['id'])[k]==v for k,v in private.items())
    assert client.post('/api/account/import',files={'file':('backup.json',json.dumps(backup).encode(),'application/json')}).status_code==200
    for format in ('txt','docx','pdf'):
        raw=client.get('/api/resumes/'+version['id']+'/export/'+format).content
        if format=='pdf': text='\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)
        elif format=='docx': text='\n'.join(p.text for p in Document(io.BytesIO(raw)).paragraphs)
        else: text=raw.decode()
        assert 'sponsorship' not in text and 'authorized_us' not in text


def test_comparison_profile_context_and_score(client,sample,account,monkeypatch):
    client.put('/api/settings',json={'authorized_us':'yes','sponsorship_now':'no'})
    monkeypatch.setattr(config,'OPENAI_KEY','fixture-no-network')
    output=fixture_analysis(account['id'],sample['data'])
    output['requirements'].append({'keyword':'Work authorization','kind':'eligibility','importance':'required','status':'supported','evidence':[{'record_id':'profile:authorized_us','quote':'authorized_us = yes'}],'job_quote':'Use SQL and Excel to prepare reports.','explanation':'Self-reported authorization.','question':''})
    calls=[]
    async def fake(body):
        calls.append(body)
        return {'model':body['model'],'output':[{'content':[{'type':'output_text','text':json.dumps(comparison_wire(output,json.loads(body['input'])['posting'],json.loads(body['input'])))}]}]}
    monkeypatch.setattr(ai,'call_api',fake)
    task=wait_task(client,client.post('/api/analyses',json={'job_id':sample['id']}))
    assert task['status']=='completed',task
    result=task['result']['data']
    assert result['score']['value']==100 and result['score']['coverage']['total']==1
    assert len(result['job_details'])==5 and all(d['value']=='Not stated' for d in result['job_details'])
    payload=json.loads(calls[0]['input'])
    assert any(s['record_id']=='profile:authorized_us' and s['quote']=='authorized_us = yes' for s in payload['private_eligibility_catalog'])
    assert 'email' not in json.dumps(payload['private_eligibility_catalog'])
    assert 'Responsibilities are not automatically prerequisites' in payload['task']
    assert 'Self-reported skills count' in payload['task']


def test_invalid_detail_and_recommendation_fail_closed(client,sample,account):
    output=fixture_analysis(account['id'],sample['data'])
    output['job_details']=[{'field':'location','value':'New York','job_quote':'invented location'}]
    with pytest.raises(ServiceError):comparison.validate(output,sample['data'],db.retrieve(account['id'],'SQL'),{})
    output=fixture_analysis(account['id'],sample['data'])
    output['recommendations']=[{'record_id':'foreign','reason':'Recommended','evidence':[]}]
    with pytest.raises(ServiceError):comparison.validate(output,sample['data'],db.retrieve(account['id'],'SQL'),{})
    output=fixture_analysis(account['id'],sample['data']);output['requirements']*=2;output['suggestions']=[]
    comparison.validate(output,sample['data'],db.retrieve(account['id'],'SQL'),{})
    assert len(output['requirements'])==1


def save(client,job,ids):
    response=client.post('/api/resumes',json={'job_id':job['id'],'experience_ids':ids})
    assert response.status_code==200,response.text
    return response.json()


def test_auto_tracker_versions_and_submitted_snapshot(client,sample,account):
    ids=[r['id'] for r in without_setup(db.list_entities(account['id'],'experience'))]
    first=save(client,sample,ids)
    entries=client.get('/api/applications').json();assert len(entries)==1
    entry=entries[0];assert entry['data']['status']=='preparing' and entry['data']['applied_date']==''
    updated={**entry['data'],'status':'interviewing','applied_date':'2026-10-20','submitted_resume_id':first['id'],'history':[{'status':'applied','date':'2026-10-20','note':'Submitted online'},{'status':'interviewing','date':'','note':''}],'interviews':[{'date':'2026-10-21','type':'Phone','outcome':'Next round'},{'date':'','type':'Technical','outcome':'Scheduled'}]}
    assert client.put('/api/applications/'+entry['id'],json=updated).status_code==200
    second=save(client,sample,ids)
    entries=client.get('/api/applications').json();assert len(entries)==1
    assert set(entries[0]['data']['resume_ids'])=={first['id'],second['id']}
    assert entries[0]['data']['submitted_resume_id']==first['id'] and entries[0]['data']['status']=='interviewing'
    db.save(account['id'],'job',{**sample['data'],'description':'Changed later posting text.'},sample['id'])
    assert entries[0]['data']['posting_snapshot']['description']==sample['data']['description']
    letter=db.save(account['id'],'cover_letter',{'job_id':sample['id'],'resume_id':first['id'],'title':'Saved cover letter','paragraphs':['Preserve this letter.']})
    current=client.get('/api/applications').json()[0]
    linked={**current['data'],'attachments':[{'kind':'cover_letter','record_id':letter['id']}]}
    for archived in (True,False,True):
        assert client.put('/api/applications/'+entry['id'],json={**linked,'archived':archived}).status_code==200
    preserved=[db.get(account['id'],id) for id in (sample['id'],first['id'],second['id'],letter['id'])]
    assert client.delete('/api/applications/'+entry['id']).status_code==200
    assert [db.get(account['id'],r['id']) for r in preserved]==preserved
    assert client.get('/api/applications').json()==[]


def test_migration_idempotent_and_restore_links(client,sample,account):
    item=without_setup(db.list_entities(account['id'],'experience'))[0]
    analysis=db.save(account['id'],'analysis',{'job_id':sample['id']})
    old=db.save(account['id'],'resume',{'analysis_id':analysis['id'],'title':'Historical','sections':[],'experience_ids':[item['id']]})
    assert len(client.get('/api/applications').json())==1
    assert len(client.get('/api/applications').json())==1
    assert db.get(account['id'],old['id'])['data']['title']=='Historical'
    backup=client.get('/api/account/export').json()
    assert client.post('/api/account/import',files={'file':('backup.json',json.dumps(backup).encode(),'application/json')}).status_code==200
    apps=client.get('/api/applications').json();assert len(apps)==2
    assert apps[0]['data']['job_id']!=apps[1]['data']['job_id']
    for a in apps:
        assert db.get(account['id'],a['data']['resume_ids'][0],'resume')
        assert db.get(account['id'],a['data']['job_id'],'job')


def test_application_isolation_dates_and_archive(client,sample,account):
    response=client.post('/api/applications',json={'title':'Outside application','status':'applied','applied_date':'2026-10-20'})
    assert response.status_code==200,response.text
    entry=response.json()
    assert client.post('/api/applications',json={'title':'Bad date','applied_date':'10/20/2026'}).status_code==422
    assert client.post('/api/applications',json={'title':'Bad date','interviews':[{'date':'2026-02-30'}]}).status_code==422
    other=TestClient(app,base_url='http://localhost')
    me=other.post('/api/auth/register',json={'username':'another','password':'another-fictional-password'}).json();other.headers['X-CSRF-Token']=me['csrf']
    onboard_test_account(me)
    assert other.get('/api/applications').json()==[]
    assert other.put('/api/applications/'+entry['id'],json=entry['data']).status_code==404
    assert other.delete('/api/applications/'+entry['id']).status_code==404
    assert other.post('/api/applications',json={'title':'Foreign posting','job_id':sample['id']}).status_code==400
    assert client.put('/api/applications/'+entry['id'],json={**entry['data'],'archived':True}).status_code==200
    assert client.get('/api/applications/dashboard').json()['counts']['submitted']==0
    assert client.get('/api/applications/dashboard?start=bad').status_code==422
    assert client.get('/api/applications/dashboard?start=2026-10-21&end=2026-10-20').status_code==422


def test_dashboard_conservation_rounds_skips_and_dates():
    def row(i,status,day='2026-10-20',history=None,interviews=None):
        return {'id':str(i),'data':{'status':status,'applied_date':day,'history':history or [],'interviews':interviews or []}}
    rows=[row(1,'rejected'),row(2,'offer',history=[{'status':'interviewing'}],interviews=[{},{}]),row(3,'screening'),row(4,'applied',''),row(5,'preparing'),row(6,'accepted',day='2026-09-01')]
    value=applications.dashboard(rows,'2026-10-01','2026-10-31')
    assert value['cohort_count']==4 and value['undated_ids']==['4']
    assert next(n['count'] for n in value['nodes'] if n['id']=='applied')==3
    assert next(n['count'] for n in value['nodes'] if n['id']=='interviewing')==1
    assert sum(e['value'] for e in value['links'] if e['source']=='applied')==3
    for n in value['nodes']:
        incoming=sum(e['value'] for e in value['links'] if e['target']==n['id'])
        outgoing=sum(e['value'] for e in value['links'] if e['source']==n['id'])
        if incoming and outgoing:assert incoming==outgoing
    assert [b for b in value['submissions'] if b['count']]==[{'date':'2026-10-19','count':4}]
    assert len(value['submissions'])==5 and value['paths']['5']==['preparing']
    assert applications.dashboard(rows,period='month')['submissions'][0]['date']=='2026-09-01'


def test_one_page_overflow_fit_and_exports(client,sample,account):
    ids=[]
    for i in range(5):
        r=client.post('/api/experience',json={'category':'project','title':'Fictional data project '+str(i),'confirmed':True,'description':'\n'.join('Analyzed SQL datasets and created documented charts for class research. '+('Details about the data and approach. '*8)+str(j) for j in range(5))}).json()
        ids.append(r['id'])
    body={'job_id':sample['id'],'experience_ids':ids}
    assert client.post('/api/resumes',json=body).status_code==400
    for template in ('education','experience'):
        body['template']=template
        task=wait_task(client,client.post('/api/resumes/fit',json=body));assert task['status']=='completed',task
        fit=task['result'];assert fit['pages']==1 and fit['omitted']
        saved=wait_task(client,client.post('/api/resumes/save',json={**body,'experience_ids':fit['experience_ids'],'content':fit['content']}));assert saved['status']=='completed',saved
        version=saved['result']
        pdf=client.get('/api/resumes/'+version['id']+'/preview')
        assert len(PdfReader(io.BytesIO(pdf.content)).pages)==1 and pdf.headers['x-frame-options']=='SAMEORIGIN'
        word=client.get('/api/resumes/'+version['id']+'/export/docx').content
        assert documents.docx_pages(word)==1
        doc=Document(io.BytesIO(word));assert doc.styles['Normal'].font.size.pt==10.5
        assert 'Analyzed SQL datasets' in client.get('/api/resumes/'+version['id']+'/export/txt').text
    assert len(without_setup(db.list_entities(account['id'],'experience')))==8


def test_failed_fit_retains_inputs(client,sample,account,monkeypatch):
    def fail(raw): raise ServiceError('Converter unavailable for test')
    monkeypatch.setattr(documents,'docx_pdf',fail)
    ids=[r['id'] for r in without_setup(db.list_entities(account['id'],'experience'))]
    task=wait_task(client,client.post('/api/resumes/fit',json={'experience_ids':ids,'job_id':sample['id']}))
    assert task['status']=='failed' and task['result']['inputs']['experience_ids']==ids
    assert len(without_setup(db.list_entities(account['id'],'experience')))==3
