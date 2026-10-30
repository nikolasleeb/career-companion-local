from conftest import without_setup
from conftest import onboard_test_account
import asyncio
import base64
import copy
import io
import json
import pytest
from docx import Document
from fastapi import HTTPException
from app import ai, db, documents, resume_conversation as guided, resume_layout, optimization
from app.models import ConversationStart, ConversationMessage, FactApproval, WorkingRender, ResumeRequest, RevisionAction, FinalizeRevision
from app.network import ServiceError
from test_optimizer import setup, fake_draft

META={'model':'test','estimated_token_cost_usd':.001,'usage':{'total_tokens':100},'reasoning_effort':'low'}

@pytest.fixture
def env(client,account,monkeypatch):
    job,rows,body=setup(client,account)
    session=guided.start(account['id'],ConversationStart(job_id=job['id']))
    calls=[]
    async def model(schema,payload,mode='economy',**kwargs):
        calls.append((schema.__name__,payload))
        if schema.__name__=='DraftAudit':return {'approved':True,'issues':[]},META
        if schema.__name__=='OptimizationOutput':
            if 'catalog' in payload:return fake_draft(payload),META
            return payload['proposal'],META
        raise AssertionError(schema)
    monkeypatch.setattr(ai,'structured',model)
    monkeypatch.setattr(resume_layout,'inspect',lambda *a:[])
    return account['id'],job,rows,session,calls

def build(env):
    user,job,rows,session,calls=env
    return asyncio.run(guided.build(user,session['id'],session['data']['version'],'economy'))

def request(state,**changes):
    rev=state['current']['data'];snap=rev['snapshot'];content={}
    selected={e['source_id']:e for s in snap['sections'] for e in s['entries']}
    for row in rev['catalog']:
        ids={i['item_id'] for i in selected.get(row['id'],{}).get('content_items',[])}
        content[row['id']]=[n for n,i in enumerate(row['items']) if i['item_id'] in ids]
    return ResumeRequest(job_id=snap['job_id'],experience_ids=snap['experience_ids'],content=content,headline=snap['headline'],template=snap['template'],**changes)

def test_first_draft_preview_finalize_exact_artifacts_and_tracker(env):
    user,job,rows,session,calls=env
    assert len(session['data']['questions'])<=3 and not calls
    assert guided.start(user,ConversationStart(job_id=job['id']))['id']==session['id']
    state=build(env);rev=state['current'];assert rev['data']['snapshot']['headline']=='Data Science | Python & SQL'
    assert rev['data']['snapshot']['layout']['pages']==1 and len(calls)==2
    assert db.list_entities(user,'resume')==[] and db.list_entities(user,'application')==[]
    assert state['data']['questions']==session['data']['questions'] # Skipping does not discard unanswered prompts.
    assert 'artifacts' not in json.dumps(state)
    finalized=guided.finalize(user,state['id'],FinalizeRevision(revision_id=rev['id'],expected_version=state['data']['version'],reviewed=True))
    saved=db.get(user,finalized['data']['last_saved_resume_id'],'resume')['data']
    raw=db.get(user,rev['id'],'resume_revision')['data']['snapshot']['artifacts']
    assert saved['artifacts']==raw
    assert documents.export_resume(saved,'pdf')[0]==base64.b64decode(raw['pdf'])
    assert documents.export_resume(saved,'docx')[0]==base64.b64decode(raw['docx'])
    assert saved['headline'].encode() in documents.export_resume(saved,'txt')[0]
    assert len(db.list_entities(user,'application'))==1
    assert db.list_entities(user,'application')[0]['data']['status']=='preparing'
    guided.finalize(user,state['id'],FinalizeRevision(revision_id=rev['id'],expected_version=finalized['data']['version'],reviewed=True))
    assert len(db.list_entities(user,'resume'))==1 and len(calls)==2

def test_local_headline_preview_and_restore(env):
    user,job,rows,session,calls=env;state=build(env);original=state['current']['id'];body=request(state);body.headline='Analysis and stakeholder communication'
    changed=guided.render(user,state['id'],WorkingRender(expected_version=state['data']['version'],request_token='new',resume=body))
    assert changed['current']['id']!=original and len(calls)==2
    snapshot=db.get(user,changed['current']['id'],'resume_revision')['data']['snapshot']
    word=Document(io.BytesIO(base64.b64decode(snapshot['artifacts']['docx'])))
    assert body.headline in '\n'.join(p.text for p in word.paragraphs)
    restored=guided.restore(user,state['id'],RevisionAction(revision_id=original,expected_version=changed['data']['version']))
    assert restored['current']['id']==original
    assert guided.state(user,state['id'])['current']['id']==original


def proposal(env,text='Wrote analytical reports for the project sponsors.',rid=None,category='project'):
    user,job,rows,session,calls=env
    return {'record_id':rid or '', 'category':category,'title':'Report project','fields':[{'field':'description','value':text,'user_quote':text}],'reason':'Supports written communication.','criterion_ids':[]}

def chat_mock(monkeypatch,value):
    async def chat(*args,**kwargs):return value,META
    monkeypatch.setattr(ai,'structured',chat)

def response(proposals=None,action='facts'):
    return {'message':'Review these details before I use them.','questions':[],'proposals':proposals or [],'action':action,'revision_instruction':'','target_item_ids':[],'target_record_ids':[],'change_headline':False,'change_selection':False}


def test_new_facts_need_approval_and_keep_provenance(env,monkeypatch):
    user,job,rows,session,calls=env;p=proposal(env);chat_mock(monkeypatch,response([p]))
    state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['value'],expected_version=0)))
    assert len(without_setup(db.list_entities(user,'experience')))==len(rows) and state['current'] is None
    p=state['data']['proposals'][0]
    result=guided.approve(user,session['id'],FactApproval(proposal_id=p['id'],expected_version=state['data']['version'],action='approve',update_draft=False,values={'description':'Wrote analytical reports for sponsors.'}))
    assert result['current'] is None and result['data']['profile_dirty']
    approved=result['data']['proposals'][0];record=db.get(user,approved['applied_record_id'],'experience')
    assert record['data']['confirmed'] and record['data']['description']=='Wrote analytical reports for sponsors.'
    assert approved['source_message_id']==result['data']['messages'][0]['id']
    with pytest.raises(HTTPException):guided.approve(user,session['id'],FactApproval(proposal_id=p['id'],expected_version=result['data']['version'],action='approve'))
    assert len(without_setup(db.list_entities(user,'experience')))==len(rows)+1


def test_concurrent_profile_update_is_never_overwritten(env,monkeypatch):
    user,job,rows,session,calls=env
    related=db.save(user,'experience',{'category':'employment','title':'Fictional related job','confirmed':True})
    rows[1]=db.save(user,'experience',{**rows[1]['data'],'linked_employment_ids':[related['id']]},rows[1]['id'])
    p=proposal(env,rid=rows[1]['id']);chat_mock(monkeypatch,response([p]))
    state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['value'],expected_version=0)))
    latest={**rows[1]['data'],'contribution':'Newer manual Profile details.'};db.save(user,'experience',latest,rows[1]['id'])
    with pytest.raises(HTTPException,match='409'):guided.approve(user,session['id'],FactApproval(proposal_id=state['data']['proposals'][0]['id'],expected_version=state['data']['version'],action='approve'))
    assert db.get(user,rows[1]['id'],'experience')['data']==latest
    assert guided.state(user,session['id'])['data']['proposals'][0]['status']=='pending'


def test_reject_and_invalid_proposal_keep_profile_unchanged(env,monkeypatch):
    user,job,rows,session,calls=env;p=proposal(env);chat_mock(monkeypatch,response([p]))
    state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['value'],expected_version=0)))
    state=guided.approve(user,session['id'],FactApproval(proposal_id=state['data']['proposals'][0]['id'],expected_version=state['data']['version'],action='reject',update_draft=False))
    assert not state['data']['profile_dirty'] and len(without_setup(db.list_entities(user,'experience')))==len(rows)
    p['fields'][0]['value']='Increased revenue 50%';chat_mock(monkeypatch,response([p]))
    with pytest.raises(ServiceError):asyncio.run(guided.send(user,session['id'],ConversationMessage(text='I wrote a report.',expected_version=state['data']['version'])))
    assert guided.state(user,session['id'])['data']['messages'][-1]['text']=='I wrote a report.'


def test_questions_leave_preview_and_wording_preserves_score(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);head=state['current']['id'];chat_mock(monkeypatch,response(action='question'))
    state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text='What would improve this?',expected_version=state['data']['version'])))
    assert state['current']['id']==head and not state['data']['profile_dirty']
    assert db.list_entities(user,'analysis')==[]


def test_failed_render_retains_inputs_and_blocks_old_finalize(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);head=state['current']['id'];body=request(state);body.headline='My retained headline'
    def fail(*args):raise ServiceError('Converter unavailable')
    monkeypatch.setattr(resume_layout,'render',fail)
    with pytest.raises(ServiceError):guided.render(user,state['id'],WorkingRender(expected_version=state['data']['version'],request_token='failed',resume=body))
    state=guided.state(user,state['id']);assert state['current']['id']==head and state['data']['workspace_request']['headline']==body.headline
    with pytest.raises(HTTPException):guided.finalize(user,state['id'],FinalizeRevision(revision_id=head,expected_version=state['data']['version'],reviewed=True))


def test_late_results_do_not_replace_current(env):
    user,job,rows,session,calls=env;state=build(env);head=state['current']['id'];v=state['data']['version']
    newer=guided.begin(user,state['id'],v,'newer')
    with pytest.raises(HTTPException):guided.save_revision(user,state['id'],v,state['current']['data']['snapshot'],[],'Late result','old')
    assert guided.state(user,state['id'])['current']['id']==head
    with pytest.raises(HTTPException):guided.save_revision(user,state['id'],newer['data']['version'],state['current']['data']['snapshot'],[],'Late result','old')


def test_account_isolation(env,client):
    user,job,rows,session,calls=env;state=build(env)
    assert client.get('/api/resume-conversations/'+session['id']).status_code==200
    other=client.post('/api/auth/register',json={'username':'other_student','password':'fictional-other-password'}).json();client.headers['X-CSRF-Token']=other['csrf']
    onboard_test_account(other)
    assert client.get('/api/resume-conversations/'+session['id']).status_code==404
    assert client.get('/api/resume-revisions/'+state['current']['id']+'/pages/1').status_code==404
    assert client.post('/api/resume-conversations',json={'job_id':job['id']}).status_code==404


def test_overflow_shows_all_pages_and_blocks_export(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);raw=db.get(user,state['current']['id'],'resume_revision')['data'];snapshot=raw['snapshot'];snapshot['layout'].update({'pages':2,'validated':False})
    db.save(user,'resume_revision',raw,state['current']['id'])
    assert guided.state(user,state['id'])['current']['data']['snapshot']['layout']['pages']==2
    with pytest.raises(HTTPException):guided.finalize(user,state['id'],FinalizeRevision(expected_version=state['data']['version'],revision_id=state['current']['id'],reviewed=True))
    assert db.list_entities(user,'resume')==[]


def test_ai_revision_does_not_commit_an_overflowing_pdf(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);head=state['current']['id']
    revisions=len(db.list_entities(user,'resume_revision'));drafts=len(db.list_entities(user,'resume_draft'))
    original_render=resume_layout.render
    def overflow(snapshot):
        value=original_render(snapshot);value['layout'].update({'pages':2,'validated':False});return value
    monkeypatch.setattr(resume_layout,'render',overflow)
    with pytest.raises(ServiceError,match='still exceeds one page'):
        asyncio.run(guided.build(user,state['id'],state['data']['version'],'economy'))
    assert guided.state(user,state['id'])['current']['id']==head
    assert len(db.list_entities(user,'resume_revision'))==revisions
    assert len(db.list_entities(user,'resume_draft'))==drafts


def test_onepage_chat_uses_measured_fitting_and_keeps_both_certifications(env,monkeypatch):
    user,job,rows,session,calls=env
    cert=rows[-1];db.save(user,'experience',{**cert['data'],'tools':'Example Cloud Certificate; Example AI Certificate'},cert['id'])
    state=build(env);head=state['current']['id'];original_model=ai.structured;original_render=resume_layout.render
    old=state['current']['data']['snapshot']
    original_count=sum(len(e['content_items']) for s in old['sections'] for e in s['entries'])
    async def model(schema,payload,*args,**kwargs):
        if schema.__name__=='ResumeChatOutput':return {**response(action='revise'),'message':'I will fit it to one page.','revision_instruction':'Fit to one page, keeping both certificates.','change_selection':True},META
        return await original_model(schema,payload,*args,**kwargs)
    def overflow(snapshot):
        value=original_render(snapshot)
        count=sum(len(e['content_items']) for s in snapshot['sections'] for e in s['entries'])
        if count>=original_count:value['layout'].update({'pages':2,'validated':False})
        return value
    monkeypatch.setattr(ai,'structured',model);monkeypatch.setattr(resume_layout,'render',overflow)
    state=asyncio.run(guided.send(user,state['id'],ConversationMessage(text='Keep it to one page and keep both certifications. Omit lower-priority material if necessary.',expected_version=state['data']['version'])))
    assert state['current']['id']!=head and state['data']['one_page_requested']
    saved=db.get(user,state['current']['id'],'resume_revision')['data']['snapshot']
    assert saved['layout']['pages']==1 and saved['layout']['validated']
    certs=[i['text'] for s in saved['sections'] for e in s['entries'] for i in e['content_items'] if i['field']=='certification']
    assert certs==['Example Cloud Certificate','Example AI Certificate']


def test_onepage_layout_can_shorten_previously_rewritten_narrative(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);original_model=ai.structured
    record=db.get(user,state['current']['id'],'resume_revision');snapshot=record['data']['snapshot']
    project=next(e for s in snapshot['sections'] for e in s['entries'] if e['category']=='project')
    target=project['content_items'][0]['item_id']
    old='Analyzed demand data using Python and presented findings clearly to classmates.'
    new='Analyzed demand data in Python; presented findings to classmates.'
    project['content_items'][0]['text']=old;project['bullets'][0]=old
    db.save(user,'resume_revision',record['data'],record['id'])
    original_render=resume_layout.render;layout_calls=[]
    def overflow(value):
        rendered=original_render(value)
        text=next(i['text'] for s in value['sections'] for e in s['entries'] for i in e['content_items'] if i['item_id']==target)
        if text!=new:rendered['layout'].update({'pages':2,'validated':False})
        return rendered
    async def model(schema,payload,*args,**kwargs):
        if schema.__name__=='ResumeChatOutput':return {**response(action='revise'),'message':'I will tighten the wording.','revision_instruction':'Keep it to one page.'},META
        value,meta=await original_model(schema,payload,*args,**kwargs)
        if schema.__name__=='OptimizationOutput' and 'target_item_ids' in payload:
            assert target in payload['target_item_ids']
            layout_calls.append(payload)
            for entry in value['entries']:
                for item in entry['items']:
                    if item['item_id']==target:item['text']=new
        return value,meta
    monkeypatch.setattr(ai,'structured',model);monkeypatch.setattr(resume_layout,'render',overflow)
    changed=asyncio.run(guided.send(user,state['id'],ConversationMessage(text='Keep it to one page by shortening wording.',expected_version=state['data']['version'])))
    result=db.get(user,changed['current']['id'],'resume_revision')['data']['snapshot']
    assert len(layout_calls)==1 and result['layout']['pages']==1
    assert result['experience_ids']==snapshot['experience_ids']
    before={i['item_id']:i['text'] for s in snapshot['sections'] for e in s['entries'] for i in e['content_items']}
    after={i['item_id']:i['text'] for s in result['sections'] for e in s['entries'] for i in e['content_items']}
    assert before.keys()==after.keys() and after[target]==new
    assert all(after[k]==v for k,v in before.items() if k!=target)
    for revision in (snapshot,result):
        doc=Document(io.BytesIO(base64.b64decode(revision['artifacts']['docx'])))
        assert doc.styles['Normal'].font.size.pt==10.5
        assert doc.styles['List Bullet'].paragraph_format.space_after.pt==3
        assert doc.styles['Normal'].paragraph_format.line_spacing==pytest.approx(1.03,abs=.001)
        assert doc.sections[0].left_margin.inches==.7


def test_real_pdf_wrapping_is_measured_for_wording_corrections(env):
    user,job,rows,session,calls=env;snapshot=guided.source(user,session['data'])
    entry=next(e for s in snapshot['sections'] for e in s['entries'] if e['category']=='project')
    item=entry['content_items'][0]
    item['text']='Analyzed operational demand data using Python, documented source limitations, prepared clear visual reports for the review team, and presented analytical findings and recommendations to classmates.'
    entry['bullets'][0]=item['text']
    result=resume_layout.render(snapshot)
    metrics=next(v for v in result['layout']['wrapped_items'] if v['item_id']==item['item_id'])
    assert metrics['record_id']==entry['source_id'] and metrics['line_count']>=2
    assert metrics['target_lines']==metrics['line_count']-1 and metrics['last_line_words']>0
    assert metrics['last_line'].rstrip('.') in item['text']
    assert not calls


def test_refinement_constrains_untargeted_edits_and_selections(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env)
    original=db.get(user,state['current']['id'],'resume_revision')['data']['snapshot']
    source=guided.source(user,state['data']);draft=fake_draft({'catalog':resume_builder_catalog(source)})
    entry=next(e for e in draft['entries'] if e['record_id']==rows[1]['id']);target=entry['items'][0]['item_id'];entry['items'][0]['text']='Analyzed demand data using Python; presented findings to classmates.'
    draft['headline']='Changed headline';next(e for e in draft['entries'] if e['record_id']!=rows[1]['id'] and e['items'])['items'][0]['text']='Wrong unrelated change'
    constrained=optimization.constrain(draft,source,original,{'target_item_ids':[target],'change_headline':False,'change_selection':False})
    assert constrained['headline']==original['headline']
    old={i['item_id']:i['text'] for s in original['sections'] for e in s['entries'] for i in e['content_items']}
    assert all(i['text']==old[i['item_id']] for e in constrained['entries'] for i in e['items'] if i['item_id']!=target)


def resume_builder_catalog(snapshot):
    from app.resume_builder import catalog
    return catalog(snapshot)


def test_approved_facts_refresh_comparison_on_next_build_only(env,monkeypatch):
    user,job,rows,session,calls=env;p=proposal(env,rid=rows[1]['id']);original_model=ai.structured
    chat_mock(monkeypatch,response([p]));state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['value'],expected_version=0)))
    state=guided.approve(user,state['id'],FactApproval(proposal_id=state['data']['proposals'][0]['id'],expected_version=state['data']['version'],action='approve',update_draft=False))
    assert not db.list_entities(user,'analysis') and not state['current']
    refreshed=[]
    async def analyze(user,posting,mode):
        from app import comparison,job_parser
        from app import profile_answers
        records=[{'id':r['id'],'kind':'experience','data':r['data']} for r in profile_answers.records(user)]
        assert any(p['fields'][0]['value'] in ai.experience_text(r['data']) for r in records)
        criteria=job_parser.criteria(posting);out={'requirements':[{'criterion_id':c['id'],'keyword':c['text'],'job_quote':c['source'],'kind':c['kind'],'importance':c['importance'],'status':'supported','evidence':[{'record_id':rows[1]['id'],'quote':p['fields'][0]['value']}],'explanation':'Supported by approved facts.','question':''} for c in criteria],'job_details':[],'recommendations':[],'suggestions':[]}
        comparison.validate(out,posting,records,{})
        out.update({'ai':META,'retrieved':records});refreshed.append(True);return out
    monkeypatch.setattr(ai,'analyze',analyze);monkeypatch.setattr(ai,'structured',original_model)
    state=asyncio.run(guided.build(user,state['id'],state['data']['version'],'economy'))
    assert refreshed==[True] and not state['data']['profile_dirty']
    analysis=db.get(user,state['data']['analysis_id'],'analysis');assert analysis['data']['score']['value']==100
    assert any(p['fields'][0]['value'] in i['original'] for r in state['current']['data']['catalog'] for i in r['items'])


def test_actionable_chat_automatically_changes_working_pdf(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);old=state['current']['id'];original_model=ai.structured
    async def model(schema,payload,*args,**kwargs):
        if schema.__name__=='ResumeChatOutput':return {**response(action='revise'),'message':'I will emphasize your analytical focus.','revision_instruction':'Change the headline to Python data analysis.','change_headline':True},META
        value,meta=await original_model(schema,payload,*args,**kwargs)
        if schema.__name__=='OptimizationOutput':value['headline']='Python data analysis'
        return value,meta
    monkeypatch.setattr(ai,'structured',model)
    state=asyncio.run(guided.send(user,state['id'],ConversationMessage(text='Change the headline to Python data analysis.',expected_version=state['data']['version'])))
    assert state['current']['id']!=old and state['current']['data']['snapshot']['headline']=='Python data analysis'
    saved=db.get(user,state['current']['id'],'resume_revision')['data']['snapshot'];word=Document(io.BytesIO(base64.b64decode(saved['artifacts']['docx'])))
    assert 'Python data analysis' in [p.text for p in word.paragraphs]
    assert db.list_entities(user,'analysis')==[] and db.list_entities(user,'resume')==[]


def test_private_authorization_proposal_never_enters_draft(env,monkeypatch):
    user,job,rows,session,calls=env
    p={'record_id':'','category':'authorization','title':'Work authorization','fields':[{'field':'sponsorship_future','value':'no','user_quote':'I do not need sponsorship in the future.'}],'reason':'Private eligibility context.','criterion_ids':[]}
    chat_mock(monkeypatch,response([p]));state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['user_quote'],expected_version=0)))
    state=guided.approve(user,state['id'],FactApproval(proposal_id=state['data']['proposals'][0]['id'],expected_version=state['data']['version'],action='approve',update_draft=False))
    assert db.settings(user)['sponsorship_future']=='no'
    snapshot=guided.source(user,state['data']);assert 'sponsorship_future' not in json.dumps(snapshot)
    assert len(without_setup(db.list_entities(user,'experience')))==len(rows)


def test_backup_restore_remaps_conversation_revision_and_provenance(env,client):
    user,job,rows,session,calls=env;state=build(env)
    backup=client.get('/api/account/export').json()
    assert {'resume_conversation','resume_revision'}<={r['kind'] for r in backup['records']}
    response=client.post('/api/account/import',files={'file':('backup.json',json.dumps(backup).encode(),'application/json')})
    assert response.status_code==200,response.text
    copies=[r for r in db.list_entities(user,'resume_conversation') if r['id']!=state['id']];assert len(copies)==1
    restored=guided.state(user,copies[0]['id']);assert restored['current']['id']!=state['current']['id']
    assert restored['current']['data']['conversation_id']==copies[0]['id'] and restored['data']['job_id']!=job['id']
    assert db.get(user,restored['current']['id'],'resume_revision')['data']['snapshot']['artifacts']


@pytest.mark.parametrize('template',['education','experience'])
def test_readable_layout_actual_pages_and_tab_stops(env,template):
    user,job,rows,session,calls=env;snapshot=guided.source(user,session['data'],template=template)
    education=next(e for s in snapshot['sections'] for e in s['entries'] if e['category']=='education')
    education['title']='Master of Science in Data Science and Applied Analytical Methods'
    education['organization']='Example University, School of Computing and Information'
    value=resume_layout.render(snapshot);assert value['layout']['pages']==1
    doc=Document(io.BytesIO(base64.b64decode(value['artifacts']['docx'])))
    assert doc.styles['Normal'].font.size.pt>=10.5
    paragraphs=[p.text for p in doc.paragraphs]
    assert not any(' | ' in p for p in paragraphs if 'Example University' in p)
    assert any('Example University' in p for p in paragraphs)
    assert documents.pdf_preview(base64.b64decode(value['artifacts']['pdf']),1).startswith(b'\x89PNG')


def test_authenticated_task_loop_local_build_render_save(env,client):
    from test_applications_workflow import wait_task
    user,job,rows,session,calls=env
    task=wait_task(client,client.post('/api/resume-conversations/'+session['id']+'/build?manual=true',json={'expected_version':0}))
    assert task['status']=='completed',task
    state=task['result'];assert state['current'] and not calls
    body=request(state);body.headline='Local headline visible before saving'
    task=wait_task(client,client.post('/api/resume-conversations/'+session['id']+'/render',json={'expected_version':state['data']['version'],'request_token':'api-edit','resume':body.model_dump()}))
    assert task['status']=='completed',task
    state=task['result'];revision=state['current']
    assert client.get('/api/resume-revisions/'+revision['id']+'/pages/1').content.startswith(b'\x89PNG')
    assert client.get('/api/resume-revisions/'+revision['id']+'/pages/2').status_code==404
    response=client.post('/api/resume-conversations/'+session['id']+'/finalize',json={'expected_version':state['data']['version'],'revision_id':revision['id'],'reviewed':True})
    assert response.status_code==200,response.text
    assert len(db.list_entities(user,'resume'))==1 and not calls


def test_unsupported_generated_claim_does_not_replace_working_revision(env,monkeypatch):
    user,job,rows,session,calls=env;state=build(env);head=state['current']['id'];original=ai.structured
    async def invalid(schema,payload,*args,**kwargs):
        value,meta=await original(schema,payload,*args,**kwargs)
        if schema.__name__=='DraftAudit':return {'approved':False,'issues':['Unsupported achievement.']},meta
        return value,meta
    monkeypatch.setattr(ai,'structured',invalid)
    with pytest.raises(ServiceError):asyncio.run(guided.build(user,state['id'],state['data']['version'],'economy'))
    assert guided.state(user,state['id'])['current']['id']==head
    assert db.list_entities(user,'resume')==[]


def test_restored_pending_proposal_can_be_approved(env,client,monkeypatch):
    user,job,rows,session,calls=env
    related=db.save(user,'experience',{'category':'employment','title':'Fictional related job','confirmed':True})
    rows[1]=db.save(user,'experience',{**rows[1]['data'],'linked_employment_ids':[related['id']]},rows[1]['id'])
    p=proposal(env,rid=rows[1]['id']);chat_mock(monkeypatch,response([p]))
    state=asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['value'],expected_version=0)))
    backup=client.get('/api/account/export').json();assert client.post('/api/account/import',files={'file':('backup.json',json.dumps(backup).encode())}).status_code==200
    restored=next(r for r in db.list_entities(user,'resume_conversation') if r['id']!=session['id'])
    facts=restored['data']['proposals'][0]
    result=guided.approve(user,restored['id'],FactApproval(expected_version=restored['data']['version'],proposal_id=facts['id'],action='approve',update_draft=False))
    assert result['data']['proposals'][0]['status']=='approved'
    assert db.get(user,rows[1]['id'],'experience')['data']==rows[1]['data']


def test_public_authorization_fact_is_rejected(env,monkeypatch):
    user,job,rows,session,calls=env;p=proposal(env,text='I am authorized to work in the United States.');chat_mock(monkeypatch,response([p]))
    with pytest.raises(ServiceError,match='private Profile'):asyncio.run(guided.send(user,session['id'],ConversationMessage(text=p['fields'][0]['value'],expected_version=0)))
    assert len(without_setup(db.list_entities(user,'experience')))==len(rows)


def test_historical_draft_opens_without_provider_or_changing_saved_versions(env):
    from app.models import LegacyDraftLoad
    user,job,rows,session,calls=env;snapshot=guided.source(user,session['data'])
    old=asyncio.run(optimization.create(user,snapshot,None,'economy'))
    count=len(calls);state=guided.adopt(user,session['id'],LegacyDraftLoad(expected_version=0,draft_id=old['id']))
    assert state['current']['data']['snapshot']['headline']==old['data']['headline']
    assert len(calls)==count and db.list_entities(user,'resume')==[]
    assert db.get(user,old['id'],'resume_draft')['data']==old['data']


def test_one_optional_layout_wording_pass_then_final_audit(env,monkeypatch):
    user,job,rows,session,calls=env
    def dangling(pdf,snapshot):
        entry=next(e for s in snapshot['sections'] for e in s['entries'] if e['category']=='project');item=entry['content_items'][0]
        return [{'kind':'dangling_words','item_id':item['item_id'],'record_id':entry['source_id'],'text':'A short carryover line.'}]
    monkeypatch.setattr(resume_layout,'inspect',dangling)
    state=build(env)
    assert [name for name,p in calls]==['OptimizationOutput','OptimizationOutput','DraftAudit']
    assert state['current']['data']['usage']['layout_ai'] and state['current']['data']['snapshot']['layout']['pages']==1


def test_real_pdf_line_inspection_detects_short_carryover():
    import shutil
    if not shutil.which('pdftotext'):pytest.skip('Real PDF text inspection is exercised in the Docker suite.')
    from reportlab.pdfgen import canvas
    buffer=io.BytesIO();page=canvas.Canvas(buffer)
    first='Reviewed the data and prepared clear analytical findings for project'
    page.drawString(50,700,first);page.drawString(50,685,'sponsors');page.save()
    text=first+' sponsors'
    snapshot={'sections':[{'entries':[{'source_id':'project','title':'Analytical project','bullets':[text],'content_items':[{'item_id':'bullet'}]}]}]}
    issues=resume_layout.inspect(buffer.getvalue(),snapshot)
    assert len(issues)==1 and issues[0]['kind']=='dangling_words' and issues[0]['ending']=='sponsors'

# Regression: choosing wording before the first draft used to retain an empty render.
def empty_workspace(env,**changes):
    user,job,rows,session,calls=env
    source=guided.source(user,session['data']);catalog=guided.resume_builder.catalog(source)
    item=next(i for r in catalog if r['category']=='project' for i in r['items'] if i['field']=='description')
    body=ResumeRequest(job_id=job['id'],experience_ids=[],content={r['id']:[] for r in catalog},headline='Python data analysis',template='experience',overrides={item['item_id']:'Analyzed demand data using Python and shared findings with classmates.'},**changes)
    return body,item

def retain_empty(env,body):
    user,job,rows,session,calls=env
    with pytest.raises(ServiceError,match='Build first draft') as failure:
        guided.render(user,session['id'],WorkingRender(expected_version=session['data']['version'],request_token='old-empty',resume=body))
    assert 'previous preview' not in str(failure.value)
    return guided.state(user,session['id'])

def test_empty_preparation_recovers_after_reload_without_provider_dispatch(env):
    user,job,rows,session,calls=env;body,item=empty_workspace(env);state=retain_empty(env,body)
    assert guided.empty_preparation(state['data']) and not calls
    state=guided.start(user,ConversationStart(job_id=job['id']))
    result=asyncio.run(guided.build(user,state['id'],state['data']['version'],'economy',manual=True))
    snap=result['current']['data']['snapshot']
    assert snap['experience_ids'] and snap['template']=='experience' and snap['headline']==body.headline
    assert any(i['text']==body.overrides[item['item_id']] for s in snap['sections'] for e in s['entries'] for i in e['content_items'])
    assert 'workspace_request' not in result['data'] and not calls

def test_empty_preparation_ai_build_keeps_edits_and_does_not_retry_provider(env):
    user,job,rows,session,calls=env;body,item=empty_workspace(env);state=retain_empty(env,body)
    result=asyncio.run(guided.build(user,state['id'],state['data']['version'],'economy'))
    assert len(calls)==2 # One draft and one audit, no comparison or retry.
    assert result['current']['data']['snapshot']['headline']==body.headline
    assert calls[0][1]['accepted_edits'][item['item_id']]==body.overrides[item['item_id']]
    assert 'workspace_request' not in result['data']

def test_failed_first_build_keeps_preparation_for_explicit_retry(env,monkeypatch):
    user,job,rows,session,calls=env;body,item=empty_workspace(env)
    original=optimization.create
    async def failed(*a,**kw):raise ServiceError('Mock provider unavailable')
    monkeypatch.setattr(optimization,'create',failed)
    with pytest.raises(ServiceError,match='Mock provider'):
        asyncio.run(guided.build(user,session['id'],session['data']['version'],'economy',preparation=body))
    saved=guided.state(user,session['id']);assert not saved['current'] and saved['data']['workspace_request']==body.model_dump()
    monkeypatch.setattr(optimization,'create',original)
    result=asyncio.run(guided.build(user,session['id'],saved['data']['version'],'economy'))
    assert result['current'] and len(calls)==2

@pytest.mark.parametrize('change',['foreign_job','foreign_analysis','unknown_override','selection'])
def test_first_build_preparation_validates_before_provider(env,change):
    user,job,rows,session,calls=env;body,item=empty_workspace(env)
    if change=='foreign_job':body.job_id='foreign'
    elif change=='foreign_analysis':body.analysis_id='foreign'
    elif change=='unknown_override':body.overrides={'foreign-item':'Edited text'}
    else:body.experience_ids=[rows[0]['id']]
    with pytest.raises((HTTPException,ServiceError)):
        asyncio.run(guided.build(user,session['id'],session['data']['version'],'economy',preparation=body))
    assert not calls and guided.state(user,session['id'])['data']['version']==session['data']['version']

def test_nonempty_retained_selection_still_requires_rendering(env):
    user,job,rows,session,calls=env;body,item=empty_workspace(env);body.experience_ids=[rows[1]['id']]
    guided.mutate(user,session['id'],session['data']['version'],lambda d,c:d.update(workspace_request=body.model_dump()))
    state=guided.state(user,session['id'])
    with pytest.raises(HTTPException,match='409'):
        asyncio.run(guided.build(user,session['id'],state['data']['version'],'economy'))
    assert not calls

def test_authenticated_first_build_accepts_preparation_once(env,client):
    from test_applications_workflow import wait_task
    user,job,rows,session,calls=env;body,item=empty_workspace(env)
    task=wait_task(client,client.post('/api/resume-conversations/'+session['id']+'/build?manual=true',json={'expected_version':session['data']['version'],'preparation':body.model_dump()}))
    assert task['status']=='completed',task
    state=task['result'];assert state['current'] and not calls
    assert state['current']['data']['snapshot']['headline']==body.headline
    with db.connect() as c:
        assert c.execute("SELECT count(*) FROM tasks WHERE kind='resume_guided' AND user_id=?",(user,)).fetchone()[0]==1
    before=state['current']['id'];state=guided.start(user,ConversationStart(job_id=job['id']))
    assert state['current']['id']==before and not calls


def test_recovered_suggestion_decisions_survive_first_build_and_local_undo(env):
    from app.models import Decision
    user,job,rows,session,calls=env
    original=rows[1]['data']['description'];proposed='Analyzed demand data with Python and shared findings with classmates.'
    analysis=db.save(user,'analysis',{'job_id':job['id'],'suggestions':[{'record_id':rows[1]['id'],'original':original,'proposed':proposed}]})
    state=guided.mutate(user,session['id'],session['data']['version'],lambda d,c:d.update(analysis_id=analysis['id']))
    body=ResumeRequest(job_id=job['id'],analysis_id=analysis['id'],experience_ids=[],content={},decisions=[Decision(index=0,action='accept',text=proposed)])
    state=asyncio.run(guided.build(user,session['id'],state['data']['version'],'economy',manual=True,preparation=body))
    assert state['current']['data']['snapshot']['decisions']==[{'index':0,'action':'accept','text':proposed}]
    item=next(i for r in state['current']['data']['catalog'] if r['id']==rows[1]['id'] for i in r['items'] if i['field']=='description')
    edited=request(state,decisions=[Decision(index=0,action='reject',text=proposed)],overrides={item['item_id']:original})
    state=guided.render(user,session['id'],WorkingRender(expected_version=state['data']['version'],request_token='undo-wording',resume=edited))
    assert state['current']['data']['snapshot']['decisions'][0]['action']=='reject'
    assert any(i['text']==original for r in state['current']['data']['catalog'] for i in r['items'])
    assert not calls


def test_refine_preserves_valid_original_headline_citation_after_checked_rewrite(env,monkeypatch):
    user,job,rows,session,calls=env
    original_model=ai.structured
    async def model(schema,payload,*args,**kwargs):
        value,meta=await original_model(schema,payload,*args,**kwargs)
        if schema.__name__=='OptimizationOutput' and 'catalog' in payload:
            for entry in value['entries']:
                for item in entry['items']:
                    if item['text']=='Analyzed demand data using Python and presented findings to classmates.':item['text']='Used Python to analyze demand data and presented findings to classmates.'
        return value,meta
    monkeypatch.setattr(ai,'structured',model)
    first=build(env)
    assert any(i['text']=='Used Python to analyze demand data and presented findings to classmates.' for s in first['current']['data']['snapshot']['sections'] for e in s['entries'] for i in e['content_items'])
    second=asyncio.run(guided.build(user,first['id'],first['data']['version'],'economy'))
    assert second['current']['id']!=first['current']['id']
    assert second['current']['data']['snapshot']['optimization']['source_checked']


def test_failed_chat_revision_retains_feedback_and_retries_without_another_chat(env,monkeypatch):
    user,job,rows,session,calls=env;first=build(env);old=first['current']['id']
    artifacts=copy.deepcopy(db.get(user,old,'resume_revision')['data']['snapshot']['artifacts'])
    original_model=ai.structured;original_create=optimization.create;chat_calls=[]
    async def model(schema,payload,*args,**kwargs):
        if schema.__name__=='ResumeChatOutput':
            chat_calls.append(True)
            return {**response(action='revise'),'message':'I will clarify the headline.','revision_instruction':'Use Python reporting as the headline.','change_headline':True},META
        result,meta=await original_model(schema,payload,*args,**kwargs)
        if schema.__name__=='OptimizationOutput':result['headline']='Python reporting'
        return result,meta
    async def fail(*args,**kwargs):raise ServiceError('Source verification failed.')
    monkeypatch.setattr(ai,'structured',model);monkeypatch.setattr(optimization,'create',fail)
    with pytest.raises(ServiceError,match='Source verification'):
        asyncio.run(guided.send(user,first['id'],ConversationMessage(text='Clarify the headline.',expected_version=first['data']['version'])))
    retained=guided.state(user,first['id']);assert retained['current']['id']==old
    assert db.get(user,old,'resume_revision')['data']['snapshot']['artifacts']==artifacts
    assert retained['data']['messages'][-1]['revision_status']=='failed'
    assert retained['data']['pending_feedback']['change_headline'] is True
    assert retained['data']['last_message_input']=='Clarify the headline.'
    monkeypatch.setattr(optimization,'create',original_create)
    updated=asyncio.run(guided.build(user,first['id'],retained['data']['version'],'economy'))
    assert updated['current']['id']!=old and updated['current']['data']['snapshot']['headline']=='Python reporting'
    assert any(m.get('revision_status')=='applied' for m in updated['data']['messages'])
    assert not updated['data'].get('pending_feedback') and not updated['data'].get('last_message_input')
    assert len(chat_calls)==1


def test_legacy_failed_chat_shows_unapplied_reply_without_rewriting_history(env):
    user,job,rows,session,calls=env;first=build(env)
    def append(data,c):
        data['workspace_token']='failed-preview'
        data['messages'] += [{'id':'request','role':'user','text':'Clarify this draft.'},{'id':'reply','role':'assistant','text':'I will clarify it.'}]
    saved=guided.mutate(user,first['id'],first['data']['version'],append)
    with db.connect() as c:
        stamp=db.now();c.execute('INSERT INTO tasks VALUES(?,?,?,?,?,?,?,?)',('failed-chat',user,'resume_guided','failed',json.dumps({'inputs':{'conversation_id':first['id'],'text':'Clarify this draft.'}}),'Source could not be verified.',stamp,stamp))
    value=guided.state(user,first['id'])
    assert value['data']['messages'][-1]['revision_status']=='failed'
    assert value['data']['last_message_input']=='Clarify this draft.'
    raw=db.get(user,first['id'],'resume_conversation')['data']
    assert 'revision_status' not in raw['messages'][-1] and 'last_message_input' not in raw
    assert value['current']['id']==first['current']['id']


def test_resume_ai_gets_only_the_approved_work_style_summary(env):
    from app import work_preferences
    user,job,rows,session,calls=env
    record=db.save(user,'work_style',{'status':'completed','report':{'domains':[{'name':'SECRET_REPORT_NAME','mean':4.75}],'facets':[],'preferences':[]},'answers':{'private':'SECRET_RAW_ANSWER'}})
    work_preferences.put(work_preferences.Change(expected_revision=0,summary='I prefer focused preparation with planned check-ins.',use_for_ai=True,reviewed=True,coach_report_access=True,assessment_id=record['id']),{'id':user})
    result=build(env)
    payload=next(value for schema,value in calls if schema=='OptimizationOutput')
    assert payload['work_style_guidance']['summary']=='I prefer focused preparation with planned check-ins.'
    assert 'SECRET_REPORT_NAME' not in json.dumps(calls) and 'SECRET_RAW_ANSWER' not in json.dumps(calls)
    assert 'professional evidence' in payload['work_style_guidance']['purpose']
    assert result['current']
