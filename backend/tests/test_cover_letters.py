from conftest import onboard_test_account
import asyncio
import base64
import copy
import io
import json
import pytest
from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader
from app import ai,config,db,documents,cover_letters as letters,applications
from app.models import CoverStart,CoverContent,CoverGenerate,CoverEdit,FinalizeRevision,RevisionAction
from app.network import ServiceError
from test_applications_workflow import wait_task

META={'model':'test','mode':'economy','reasoning_effort':'low','estimated_token_cost_usd':.001,'usage':{'total_tokens':100}}

@pytest.fixture
def env(client,account,monkeypatch):
    job=client.post('/api/jobs',json={'title':'Analyst Intern','company':'Fictional Employer','description':'Analyst Intern\nRequired qualifications\nAnalytical project experience.\nPreferred qualifications\nPython or SQL.\nWork authorization\nNo sponsorship.'}).json()
    source={'schema_version':4,'job_id':job['id'],'posting_snapshot':job['data'],'title':'Old resume name','name':'Fictional Student','headline':'Data Analysis','contact_lines':['student@example.com'],'sections':[{'title':'Projects','entries':[{'category':'project','source_id':'original-profile-record','title':'Demand analysis','organization':'Example University','details':'Example University | 2026','dates':'2026','bullets':['Analyzed demand data using Python and SQL.','Wrote a report explaining the findings for classmates.'],'description':'OMITTED PRIVATE FULL RECORD TEXT: used Java and increased revenue 50%.','content_items':[{'original':'SECRET OMITTED ITEM','text':'SECRET OMITTED ITEM'}]}]}],'layout':{'pages':1,'validated':True}}
    resume=db.save(account['id'],'resume',source);applications.attach_resume(account['id'],resume)
    db.update_settings(account['id'],{'authorized_us':'yes','sponsorship_now':'no','sponsorship_future':'no'})
    state=letters.start(account['id'],CoverStart(job_id=job['id'],resume_id=resume['id'],letter_date='2026-10-21'))
    calls=[]
    async def model(schema,payload,*args,**kwargs):
        calls.append((schema.__name__,payload))
        if schema.__name__=='DraftAudit':return {'approved':True,'issues':[]},META
        rows=payload['saved_resume_sources'];body=[r for r in rows if r['id'].endswith((':0',':1'))]
        return {'paragraphs':[{'text':'I am interested in the Analyst Intern position at Fictional Employer.','evidence':[{'source_id':'posting','quote':'Analyst Intern'}]},{'text':body[0]['text'],'evidence':[{'source_id':body[0]['id'],'quote':body[0]['text']}]},{'text':body[1]['text'],'evidence':[{'source_id':body[1]['id'],'quote':body[1]['text']}]}],'summary':'Drafted a letter using your analytical project.'},META
    monkeypatch.setattr(ai,'structured',model);monkeypatch.setattr(config,'OPENAI_KEY','fictional-no-network')
    return account['id'],job,resume,state,calls


def generate(env):
    user,job,resume,state,calls=env
    return asyncio.run(letters.generate(user,state['id'],CoverGenerate(expected_version=state['data']['version'],content=CoverContent.model_validate(state['data']['content']))))


def test_exact_saved_resume_sources_and_privacy(env):
    user,job,resume,state,calls=env
    assert letters.start(user,CoverStart(job_id=job['id'],resume_id=resume['id']))['id']==state['id']
    assert not calls and letters.estimate(user,state['id'],'economy')['usd']>0
    result=generate(env)
    assert len(calls)==2 and result['current']['data']['layout']['pages']==1
    request=json.dumps(calls)
    for secret in ('OMITTED PRIVATE','SECRET OMITTED','sponsorship_now','authorized_us','increased revenue 50%'):
        assert secret not in request
    assert 'using Python and SQL' in request
    assert not db.list_entities(user,'cover_letter') and 'artifacts' not in json.dumps(result)
    # A later Profile/source-resume change cannot replace the workspace's exact source snapshot.
    newer=copy.deepcopy(resume['data']);newer['sections'][0]['entries'][0]['bullets']=['CHANGED LATER SOURCE']
    db.save(user,'resume',newer,resume['id'])
    restored=letters.state(user,state['id']);assert 'CHANGED LATER SOURCE' not in json.dumps(restored['data']['resume_snapshot'])


def test_review_finalize_same_artifacts_filenames_and_application_status(env):
    user,job,resume,state,calls=env;result=generate(env);head=result['current'];raw=db.get(user,head['id'],'cover_revision')['data']
    application=db.list_entities(user,'application')[0];tracked={**application['data'],'status':'interviewing','submitted_resume_id':resume['id'],'history':[{'status':'applied','date':'2026-10-01','note':''}],'applied_date':'2026-10-01'};db.save(user,'application',tracked,application['id'])
    saved=letters.finalize(user,result['id'],FinalizeRevision(expected_version=result['data']['version'],revision_id=head['id'],reviewed=True))
    record=db.get(user,saved['data']['last_saved_letter_id'],'cover_letter');data=record['data']
    assert data['artifacts']==raw['artifacts'] and data['resume_id']==resume['id']
    for fmt in ('pdf','docx'):assert letters.export_data(data,fmt)[0]==base64.b64decode(raw['artifacts'][fmt])
    assert b'Wrote a report' in letters.export_data(data,'txt')[0]
    assert letters.filename(data,'pdf')=='Fictional-Student-Fictional-Employer-Analyst-Intern-Cover-Letter.pdf'
    app=db.get(user,application['id'],'application')['data']
    assert app['status']=='interviewing' and app['applied_date']=='2026-10-01' and app['submitted_resume_id']==resume['id']
    assert app['attachments']==[{'kind':'cover_letter','record_id':record['id']}]
    letters.finalize(user,result['id'],FinalizeRevision(expected_version=saved['data']['version'],revision_id=head['id'],reviewed=True))
    assert len(db.list_entities(user,'cover_letter'))==1 and len(calls)==2


def test_manual_edits_render_without_provider_and_restore(env):
    user,job,resume,state,calls=env
    content=CoverContent(letter_date='2026-10-21',subject='Application for Analyst Intern',paragraphs=['I analyzed demand data using Python and SQL.','I wrote a report explaining the findings for classmates.'])
    state=asyncio.run(letters.render(user,state['id'],CoverEdit(expected_version=0,content=content,request_token='local-one')))
    assert not calls;first=state['current']['id'];content.greeting='Dear Analyst Recruiting Team,'
    changed=asyncio.run(letters.render(user,state['id'],CoverEdit(expected_version=state['data']['version'],content=content,request_token='local-two')))
    raw=db.get(user,changed['current']['id'],'cover_revision')['data'];word=Document(io.BytesIO(base64.b64decode(raw['artifacts']['docx'])))
    assert content.greeting in [p.text for p in word.paragraphs] and 11<=word.styles['Normal'].font.size.pt<=12
    assert len(PdfReader(io.BytesIO(base64.b64decode(raw['artifacts']['pdf']))).pages)==1
    restored=letters.restore(user,state['id'],RevisionAction(expected_version=changed['data']['version'],revision_id=first))
    assert restored['current']['id']==first and not calls


def test_failure_keeps_latest_inputs_and_blocks_stale_save(env,monkeypatch):
    user,job,resume,state,calls=env;state=generate(env);head=state['current']['id'];content=CoverContent.model_validate(state['data']['content']);content.paragraphs[0]='My retained manual change.'
    def fail(*args):raise ServiceError('Converter unavailable')
    monkeypatch.setattr(letters,'render_document',fail)
    with pytest.raises(ServiceError):asyncio.run(letters.render(user,state['id'],CoverEdit(expected_version=state['data']['version'],content=content,request_token='failed')))
    reloaded=letters.state(user,state['id']);assert reloaded['current']['id']==head and reloaded['preview_stale']
    assert reloaded['data']['content']['paragraphs'][0]=='My retained manual change.'
    with pytest.raises(HTTPException):letters.finalize(user,state['id'],FinalizeRevision(expected_version=reloaded['data']['version'],revision_id=head,reviewed=True))


def test_source_audit_and_invalid_references_never_replace_current(env,monkeypatch):
    user,job,resume,state,calls=env;state=generate(env);head=state['current']['id'];original=ai.structured
    async def bad(schema,payload,*args,**kwargs):
        value,meta=await original(schema,payload,*args,**kwargs)
        if schema.__name__=='DraftAudit':return {'approved':False,'issues':['Unsupported claim.']},meta
        return value,meta
    monkeypatch.setattr(ai,'structured',bad)
    with pytest.raises(ServiceError,match='source check'):asyncio.run(letters.generate(user,state['id'],CoverGenerate(expected_version=state['data']['version'],content=CoverContent.model_validate(state['data']['content']))))
    assert letters.state(user,state['id'])['current']['id']==head
    with pytest.raises(ServiceError):letters.validate_output({'paragraphs':[{'text':'I increased revenue 90%.','evidence':[{'source_id':'missing','quote':'revenue'}]}]},letters.sources(state['data']),job['data'])
    with pytest.raises(ServiceError):letters.validate_output({'paragraphs':[{'text':'I increased revenue 90%.','evidence':[{'source_id':'posting','quote':'Analyst Intern'}]}]},letters.sources(state['data']),job['data'])


def test_wrong_posting_foreign_accounts_and_attachments(env,client):
    user,job,resume,state,calls=env;result=generate(env);saved=letters.finalize(user,state['id'],FinalizeRevision(expected_version=result['data']['version'],revision_id=result['current']['id'],reviewed=True));letter=db.get(user,saved['data']['last_saved_letter_id'],'cover_letter')
    otherjob=client.post('/api/jobs',json={'title':'Other role','description':'A different posting requiring Python.'}).json()
    with pytest.raises(HTTPException):letters.start(user,CoverStart(job_id=otherjob['id'],resume_id=resume['id']))
    bad={'job_id':otherjob['id'],'resume_ids':[resume['id']],'attachments':[{'kind':'cover_letter','record_id':letter['id']}]}
    with pytest.raises(ServiceError):applications.validate_links(user,bad)
    other=client.post('/api/auth/register',json={'username':'second_student','password':'fictional-other-password'}).json();client.headers['X-CSRF-Token']=other['csrf']
    onboard_test_account(other)
    assert client.get('/api/cover-workspaces/'+state['id']).status_code==404
    assert client.get('/api/cover-revisions/'+result['current']['id']+'/pages/1').status_code==404
    assert client.get('/api/cover-letters/'+letter['id']+'/export/pdf').status_code==404
    assert client.post('/api/cover-letters/workspace',json={'job_id':job['id'],'resume_id':resume['id']}).status_code==404


def test_backup_restore_remaps_exact_resume_letter_and_tracker(env,client):
    user,job,resume,state,calls=env;state=generate(env);saved=letters.finalize(user,state['id'],FinalizeRevision(expected_version=state['data']['version'],revision_id=state['current']['id'],reviewed=True))
    backup=client.get('/api/account/export').json();response=client.post('/api/account/import',files={'file':('backup.json',json.dumps(backup).encode(),'application/json')});assert response.status_code==200,response.text
    restored=next(r for r in db.list_entities(user,'cover_workspace') if r['id']!=state['id']);state=letters.state(user,restored['id'])
    assert state['data']['resume_id']!=resume['id'] and state['current']['data']['workspace_id']==state['id']
    letter=next(r for r in db.list_entities(user,'cover_letter') if r['data']['workspace_id']==state['id'])
    app=next(r for r in db.list_entities(user,'application') if r['data']['job_id']==state['data']['job_id'])
    assert app['data']['attachments']==[{'kind':'cover_letter','record_id':letter['id']}] and letter['data']['resume_id'] in app['data']['resume_ids']
    assert letter['data']['artifacts']


def test_authenticated_tasks_manual_flow_and_download_headers(env,client):
    user,job,resume,state,calls=env
    content=CoverContent(paragraphs=['I analyzed demand data using Python and SQL.'])
    task=wait_task(client,client.post('/api/cover-workspaces/'+state['id']+'/render',json={'expected_version':0,'content':content.model_dump()}));assert task['status']=='completed',task
    state=task['result'];assert client.get('/api/cover-revisions/'+state['current']['id']+'/pages/1').content.startswith(b'\x89PNG')
    response=client.post('/api/cover-workspaces/'+state['id']+'/finalize',json={'expected_version':state['data']['version'],'revision_id':state['current']['id'],'reviewed':True});assert response.status_code==200,response.text
    rid=response.json()['data']['last_saved_letter_id']
    for fmt in ('pdf','docx','txt'):
        response=client.get('/api/cover-letters/'+rid+'/export/'+fmt);assert response.status_code==200
        assert 'Fictional-Student-Fictional-Employer-Analyst-Intern-Cover-Letter.'+fmt in response.headers['content-disposition']
    assert not calls


def test_overflow_shows_multiple_pages_blocks_final_export(env):
    user,job,resume,state,calls=env
    content=CoverContent(paragraphs=[' '.join(['Detailed analytical experience and thoughtful communication with collaborators.']*30)]*4)
    state=asyncio.run(letters.render(user,state['id'],CoverEdit(expected_version=0,content=content)))
    assert state['current']['data']['layout']['pages']>1 and not state['current']['data']['layout']['validated']
    with pytest.raises(HTTPException):letters.finalize(user,state['id'],FinalizeRevision(expected_version=state['data']['version'],revision_id=state['current']['id'],reviewed=True))
    assert not db.list_entities(user,'cover_letter')


def test_late_revision_cannot_replace_newer_input(env):
    user,job,resume,state,calls=env
    first,token=letters.begin(user,state['id'],CoverEdit(expected_version=0,content=CoverContent(paragraphs=['First request.'])))
    letters.begin(user,state['id'],CoverEdit(expected_version=first['data']['version'],content=CoverContent(paragraphs=['Newer request.'])))
    with pytest.raises(HTTPException):letters.commit_revision(user,first,token,{}, {},{},[],'Late result')
    assert letters.state(user,state['id'])['data']['content']['paragraphs']==['Newer request.']


def chat_model(monkeypatch,output):
    original=ai.structured
    async def fake(schema,payload,*args,**kwargs):
        if schema.__name__=='CoverChatOutput':return {'action':'reply','reply':'Use the strongest supported example.','revision_instruction':'','paragraph_indices':[],'revise_all':False,'header_style':'unchanged','subject':'unchanged',**output},META
        return await original(schema,payload,*args,**kwargs)
    monkeypatch.setattr(ai,'structured',fake)


def test_chat_question_persists_without_changing_preview(env,monkeypatch,client):
    from app.models import CoverMessage
    user,job,resume,state,calls=env;state=generate(env);head=state['current']['id'];chat_model(monkeypatch,{})
    state=asyncio.run(letters.send(user,state['id'],CoverMessage(expected_version=state['data']['version'],content=CoverContent.model_validate(state['data']['content']),text='Which example is strongest?',request_token='question-one')))
    assert state['current']['id']==head and not state['preview_stale']
    assert state['data']['messages'][-2]['role']=='user' and state['data']['messages'][-1]['role']=='assistant'
    reloaded=client.get('/api/cover-workspaces/'+state['id']).json();assert reloaded['data']['messages']==state['data']['messages']
    assert len(calls)==2 # Only initial generation and audit in the fixture; no redrafting for the question.


def test_chat_targeted_revision_preserves_unrelated_manual_wording(env,monkeypatch):
    from app.models import CoverMessage
    user,job,resume,state,calls=env;state=generate(env);content=CoverContent.model_validate(state['data']['content'])
    content.paragraphs[2]='I wrote a report for classmates explaining the findings.'
    chat_model(monkeypatch,{'action':'revise','reply':'I clarified the second paragraph.','revision_instruction':'Clarify paragraph two. Preserve others exactly.','paragraph_indices':[1]})
    state=asyncio.run(letters.send(user,state['id'],CoverMessage(expected_version=state['data']['version'],content=content,text='Clarify paragraph 2.',request_token='revise-one')))
    assert state['current']['data']['content']['paragraphs'][2]==content.paragraphs[2]
    assert state['current']['data']['usage']['chat_ai'] and not state['preview_stale']
    assert len(calls)==4 and state['data']['messages'][-1]['role']=='assistant'


def test_chat_new_facts_and_failure_keep_sources_and_history(env,monkeypatch):
    from app.models import CoverMessage
    user,job,resume,state,calls=env;state=generate(env);head=state['current']['id']
    chat_model(monkeypatch,{'reply':'That result is absent from the selected resume. Update Profile and save a new resume first.'})
    body=CoverMessage(expected_version=state['data']['version'],content=CoverContent.model_validate(state['data']['content']),text='Add that I increased revenue 90%.',request_token='unsupported')
    state=asyncio.run(letters.send(user,state['id'],body));assert state['current']['id']==head
    assert '90%' not in json.dumps(state['current']['data']['content']) and len(calls)==2
    async def failed(*args,**kwargs):raise ServiceError('Provider unavailable')
    monkeypatch.setattr(ai,'structured',failed)
    body=body.model_copy(update={'text':'Make it warmer.','expected_version':state['data']['version'],'request_token':'retry-me'})
    with pytest.raises(ServiceError):asyncio.run(letters.send(user,state['id'],body))
    state=letters.state(user,state['id']);assert state['data']['messages'][-1]['text']=='Make it warmer.' and state['current']['id']==head
    with pytest.raises(ServiceError):asyncio.run(letters.send(user,state['id'],body.model_copy(update={'expected_version':state['data']['version']})))
    assert sum(m['text']=='Make it warmer.' for m in letters.state(user,state['id'])['data']['messages'])==1


def test_chat_header_only_renders_locally_no_body_rewrite(env,monkeypatch):
    from app.models import CoverMessage
    user,job,resume,state,calls=env;state=generate(env);before=state['current']['data']['content']['paragraphs']
    chat_model(monkeypatch,{'action':'revise','reply':'I showed the subject line.','subject':'show'})
    state=asyncio.run(letters.send(user,state['id'],CoverMessage(expected_version=state['data']['version'],content=CoverContent.model_validate(state['data']['content']),text='Show the subject line.')))
    assert state['current']['data']['content']['show_subject'] and state['current']['data']['content']['paragraphs']==before
    assert not state['preview_stale'] and len(calls)==2


def test_header_privacy_and_measured_readable_page_target(env):
    user,job,resume,state,calls=env
    snapshot=state['data']['resume_snapshot'];snapshot['contact_lines']=['student@example.com | 555-0100 | Tucson, AZ','https://github.com/example | https://linkedin.com/in/example']
    paragraph=('I analyzed demand data using Python and SQL, created charts and wrote a report explaining findings for classmates. This project gave me practice organizing information, explaining my methods, and communicating analytical findings to an audience. ')
    content=CoverContent(subject='Optional role subject',paragraphs=[paragraph*2]*4).model_dump()
    layout,artifacts=letters.render_document(snapshot,job['data'],content)
    assert layout['pages']==1 and 11<=layout['body_font_pt']<=12
    text=PdfReader(io.BytesIO(base64.b64decode(artifacts['pdf']))).pages[0].extract_text()
    assert 'github' not in text and 'Optional role subject' not in text
    assert 'student@example.com' in text and '555-0100' in text
    if letters.shutil.which('pdftotext'):assert layout['page_fill']>=2/3 and not layout['below_target']
    _,short=letters.render_document(snapshot,job['data'],{**content,'paragraphs':['I analyzed demand data using Python and SQL.']})
    assert 'https://github.com' in letters.plain(snapshot,job['data'],{**content,'header_style':'full'})
    assert 'github' not in letters.plain(snapshot,job['data'],content)


def test_chat_task_backup_recovery_and_account_scope(env,monkeypatch,client):
    user,job,resume,state,calls=env;state=generate(env);chat_model(monkeypatch,{})
    task=wait_task(client,client.post('/api/cover-workspaces/'+state['id']+'/message',json={'expected_version':state['data']['version'],'content':state['data']['content'],'text':'Which example should I emphasize?','request_token':'http-chat'}))
    assert task['status']=='completed',task
    state=task['result'];backup=client.get('/api/account/export').json()
    response=client.post('/api/account/import',files={'file':('backup.json',json.dumps(backup).encode(),'application/json')});assert response.status_code==200,response.text
    restored=next(r for r in db.list_entities(user,'cover_workspace') if r['id']!=state['id'])
    assert restored['data']['messages'][-2]['text']=='Which example should I emphasize?'
    assert restored['data']['resume_id']!=resume['id'] and letters.state(user,restored['id'])['current']
    other=client.post('/api/auth/register',json={'username':'other_chat_student','password':'fictional-other-password'}).json();client.headers['X-CSRF-Token']=other['csrf']
    onboard_test_account(other)
    response=client.post('/api/cover-workspaces/'+state['id']+'/message',json={'expected_version':state['data']['version'],'content':state['data']['content'],'text':'Change it.'})
    assert response.status_code in (404,503) # No provider access, and no record ownership, for this account.
    assert client.get('/api/cover-workspaces/'+state['id']).status_code==404


def test_exact_posting_quote_uses_original_newlines_not_json_escapes():
    posting={'title':'Analyst Intern','company':'Example Company','description':'Analyst Intern\nAnalyze demand data and explain findings.'}
    letters.validate_output({'paragraphs':[{'text':'I am applying for the Analyst Intern role.','evidence':[{'source_id':'posting','quote':'Analyst Intern\nAnalyze demand data and explain findings.'}]}]},[],posting)
    with pytest.raises(ServiceError):letters.validate_output({'paragraphs':[{'text':'I am applying.','evidence':[{'source_id':'posting','quote':'Lead a department and double revenue.'}]}]},[],posting)
