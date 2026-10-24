from conftest import comparison_wire
import asyncio
import copy
import io
import json
import socket
import time
from datetime import datetime,timezone,timedelta
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from pypdf import PdfReader
from docx import Document
from reportlab.pdfgen import canvas
from app import db, config, ai, documents, integrations
from app.main import app
from app.network import ServiceError, public_address, html_job
from app.models import SearchRequest, AnalysisOutput

def wait(client,task_id):
    for _ in range(300):
        result=client.get('/api/tasks/'+task_id).json()
        if result['status'] not in ('queued','running'): return result
        time.sleep(.01)
    raise AssertionError('Task did not complete')

def fixture_analysis(user,job):
    record=db.list_entities(user,'experience')[0]
    guide=json.loads((config.CONTENT_DIR/'guidance.json').read_text())[0]
    original=record['data']['description'].split('. ')[0]
    return {'summary':'Review your confirmed SQL and Excel experience. Python needs clarification.',
      'job_details':[], 'recommendations':[],
      'requirements':[{'kind':'qualification','keyword':'SQL','job_quote':'Use SQL and Excel to prepare reports.','importance':'required','status':'supported','evidence':[{'record_id':record['id'],'quote':original}],'explanation':'Check this excerpt.','question':''}],
      'suggestions':[{'record_id':record['id'],'original':original,'proposed':original,'reason':'Keep the verified action clear.','guide_id':guide['id'],'guide_quote':guide['text']}],
      'questions':['Have you used Python in a class or project?']}

def test_accounts_sessions_csrf_and_isolation(client,account):
    assert client.get('/api/me').json()['username']=='student'
    assert client.put('/api/settings',json={'name':'A'},headers={'X-CSRF-Token':'wrong'}).status_code==403
    assert client.put('/api/settings',json={'name':'A'},headers={'Origin':'https://unrelated.example'}).status_code==403
    item=client.post('/api/experience',json={'title':'Private project','description':'Used SQL.','confirmed':True}).json()
    other=TestClient(app,base_url='http://localhost')
    login=other.post('/api/auth/register',json={'username':'other','password':'other-fictional-password'}).json()
    other.headers['X-CSRF-Token']=login['csrf']
    assert other.get('/api/workspace').json()['experience']==[]
    assert other.put('/api/experience/'+item['id'],json={'title':'Changed'}).status_code==404
    assert other.delete('/api/records/'+item['id']).status_code==404
    assert item['id'] not in {r['id'] for r in db.retrieve(login['id'],'SQL')}
    assert client.post('/api/auth/logout').status_code==200
    assert client.get('/api/me').status_code==401
    assert client.post('/api/auth/login',json={'username':'student','password':'fictional-test-password'}).status_code==200

def test_validation_does_not_echo_password(client):
    response=client.post('/api/auth/register',json={'username':'a','password':'secret'})
    assert response.status_code==422
    assert 'secret' not in response.text

@pytest.mark.parametrize('scenario',['analyst','science','it'])
def test_fictional_flows_exports_and_master_preservation(client,account,scenario):
    client.post('/api/demo/'+scenario)
    before=client.get('/api/workspace').json()
    ids=[r['id'] for r in before['experience']]
    for template in ('education','experience'):
        response=client.post('/api/resumes',json={'title':'Classroom resume','template':template,'experience_ids':ids})
        assert response.status_code==200,response.text
        version=response.json()
        for format in ('txt','pdf','docx'):
            exported=client.get('/api/resumes/'+version['id']+'/export/'+format)
            assert exported.status_code==200,exported.text
            if format=='pdf':
                pdf=PdfReader(io.BytesIO(exported.content));text=' '.join(page.extract_text() for page in pdf.pages)
                assert len(pdf.pages)<=3
            elif format=='docx':text=' '.join(p.text for p in Document(io.BytesIO(exported.content)).paragraphs)
            else:text=exported.text
            assert version['data']['name'] in text
            assert '[missing' not in text.lower()
            assert 'fictional-test-password' not in text
    after=client.get('/api/workspace').json()
    assert before['experience']==after['experience']
    assert len(after['resume'])==2

def test_unconfirmed_and_missing_fields(client,account):
    draft=client.post('/api/experience',json={'title':'Draft internship','confirmed':False}).json()
    assert client.post('/api/resumes',json={'experience_ids':[draft['id']]}).status_code==400
    entry=client.post('/api/experience',json={'title':'Class project','description':'Used SQL.','confirmed':True,'dates':'[date missing]'}).json()
    version=client.post('/api/resumes',json={'experience_ids':[entry['id']]}).json()
    text=client.get('/api/resumes/'+version['id']+'/export/txt').text
    assert '[date missing]' not in text
    assert 'Resume' in text
    assert version['data']['warnings']

def test_text_pdf_import_boundaries(client,account):
    response=client.post('/api/documents',files={'file':('resume.txt',b'Course project: used SQL to summarize a dataset.','text/plain')})
    assert response.status_code==200,response.text
    document=response.json()
    assert not document['data']['reviewed']
    assert client.get('/api/workspace').json()['experience']==[]
    review=client.put('/api/documents/'+document['id']+'/review',json={'text':'Reviewed: I wrote the SQL query.'})
    assert review.json()['data']['reviewed']
    assert client.post('/api/documents',files={'file':('wrong.exe',b'test')}).status_code==400
    assert client.post('/api/documents',files={'file':('empty.txt',b'')}).status_code==400
    assert client.post('/api/documents',files={'file':('large.txt',b'a'*(config.MAX_UPLOAD+1))}).status_code==400
    output=io.BytesIO();pdf=canvas.Canvas(output)
    for _ in range(31):pdf.drawString(50,700,'Too many pages');pdf.showPage()
    pdf.save()
    assert client.post('/api/documents',files={'file':('long.pdf',output.getvalue())}).status_code==400
    output=io.BytesIO();pdf=canvas.Canvas(output);pdf.showPage();pdf.save()
    response=client.post('/api/documents',files={'file':('scan.pdf',output.getvalue())})
    assert response.status_code==400 and 'scanned' in response.text

def test_docx_zip_guard_and_page_limit(monkeypatch):
    output=io.BytesIO();doc=Document();doc.add_paragraph('Used SQL for a class project.');doc.save(output)
    monkeypatch.setattr(documents,'docx_pages',lambda raw:1)
    assert 'SQL' in documents.extract('resume.docx',output.getvalue())['text']
    monkeypatch.setattr(documents,'docx_pages',lambda raw:31)
    with pytest.raises(ServiceError,match='30-page'):documents.extract('resume.docx',output.getvalue())

def test_missing_integrations_are_explicit(client,sample):
    assert client.post('/api/analyses',json={'job_id':sample['id']}).status_code==503
    response=client.get('/api/assessment/questions')
    assert response.status_code==400 and 'ONET_API_KEY' in response.text
    assert client.get('/api/resources').json()

def test_structured_ai_modes_and_persistent_results(client,sample,account,monkeypatch):
    monkeypatch.setattr(config,'OPENAI_KEY','fixture-not-a-real-key')
    output=fixture_analysis(account['id'],sample['data']);calls=[]
    async def fake(body):
        calls.append(body)
        return {'status':'completed','model':config.MODEL,'usage':{'input_tokens':100,'output_tokens':50,'total_tokens':150},'output':[{'type':'message','content':[{'type':'output_text','text':json.dumps(comparison_wire(output,json.loads(body['input'])['posting'],json.loads(body['input'])))}]}]}
    monkeypatch.setattr(ai,'call_api',fake)
    client.put('/api/settings',json={'mode':'deep'})
    assert calls==[]
    for mode,effort in [('economy','low'),('deep','high')]:
        response=client.post('/api/analyses',json={'job_id':sample['id'],'mode':mode})
        assert response.status_code==200,response.text
        result=wait(client,response.json()['task_id'])
        assert result['status']=='completed',result
        data=result['result']['data'];assert data['ai']['mode']==mode and data['ai']['usage']['total_tokens']==150
        assert calls[-1]['reasoning']['effort']==effort
        assert calls[-1]['store'] is False
        assert 'untrusted' in calls[-1]['instructions']
    assert len(client.get('/api/workspace').json()['analysis'])==2

def test_bad_ai_references_and_invented_numbers(client,sample,account):
    output=fixture_analysis(account['id'],sample['data']);retrieved=db.retrieve(account['id'],'SQL')
    guides=json.loads((config.CONTENT_DIR/'guidance.json').read_text())
    retrieved += [{'id':g['id'],'kind':'guidance','data':g} for g in guides]
    ai.validate_analysis(copy.deepcopy(output),sample['data'],retrieved)
    bad=copy.deepcopy(output);bad['suggestions'][0]['proposed']='Increased revenue by 92%.'
    with pytest.raises(ServiceError,match='number'):ai.validate_analysis(bad,sample['data'],retrieved)
    bad=copy.deepcopy(output);bad['requirements'][0]['evidence'][0]['record_id']='invented'
    with pytest.raises(ServiceError,match='reference'):ai.validate_analysis(bad,sample['data'],retrieved)
    bad=copy.deepcopy(output);bad['suggestions'][0]['guide_quote']='made up citation'
    with pytest.raises(ServiceError,match='reference'):ai.validate_analysis(bad,sample['data'],retrieved)

def test_negation_and_coursework_guard():
    job={'description':'Python is required.'}
    record={'id':'r','kind':'experience','data':{'category':'coursework','description':'No Python experience.'}}
    guide={'id':'g','kind':'guidance','data':{'text':'Use supported evidence.'}}
    output={'requirements':[{'keyword':'Python','job_quote':'Python is required.','status':'supported','evidence':[{'record_id':'r','quote':'No Python experience.'}]}],'suggestions':[]}
    assert ai.validate_analysis(output,job,[record,guide])['requirements'][0]['status']=='needs_confirmation'
    output['suggestions']=[{'record_id':'r','original':'No Python experience.','proposed':'Used Python professionally.','guide_id':'g','guide_quote':'Use supported evidence.'}]
    with pytest.raises(ServiceError):ai.validate_analysis(output,job,[record,guide])

def test_invalid_ai_output_and_interrupted_tasks(client,sample,monkeypatch):
    monkeypatch.setattr(config,'OPENAI_KEY','fixture')
    async def bad(body):return {'output':[{'content':[{'type':'output_text','text':'{"wrong":true}'}]}]}
    monkeypatch.setattr(ai,'call_api',bad)
    task=client.post('/api/analyses',json={'job_id':sample['id']}).json()
    result=wait(client,task['task_id']);assert result['status']=='failed' and 'contract' in result['error']
    with db.connect() as c:c.execute("UPDATE tasks SET status='running' WHERE id=?",(task['task_id'],))
    db.init()
    assert client.get('/api/tasks/'+task['task_id']).json()['status']=='interrupted'
    assert client.get('/api/workspace').json()['job']

def test_edit_decisions_do_not_overwrite_master(client,sample,account):
    output=fixture_analysis(account['id'],sample['data'])
    analysis=db.save(account['id'],'analysis',{**output,'job_id':sample['id']})
    before=db.list_entities(account['id'],'experience')
    response=client.post('/api/resumes',json={'experience_ids':[r['id'] for r in before],'analysis_id':analysis['id'],'decisions':[{'index':0,'action':'edit','text':'Student-confirmed wording for this classroom exercise.'}]})
    assert response.status_code==200,response.text
    text=client.get('/api/resumes/'+response.json()['id']+'/export/txt').text
    assert 'Student-confirmed wording' in text
    assert before==db.list_entities(account['id'],'experience')

def test_onet_questions_progress_and_official_scoring(client,account,monkeypatch):
    monkeypatch.setattr(config,'ONET_KEY','fixture')
    calls=[]
    areas=['realistic','investigative','artistic','social','enterprising','conventional']
    async def fake(path,params=None):
        calls.append((path,params))
        if path=='questions_30':return {'question':[{'index':i,'area':areas[i%6],'text':'Fixture activity '+str(i)} for i in range(1,31)],'answer_option':[{'value':i,'name':str(i)} for i in range(1,6)]}
        if path=='results':return {'result':[{'code':a,'title':a.title(),'score':15,'description':'Provider fixture.'} for a in areas]}
        return {'career':[{'code':'15-2051.00','title':'Data Scientists','fit':'Best'}]}
    monkeypatch.setattr(integrations,'onet',fake)
    assert len(client.get('/api/assessment/questions').json()['questions'])==30
    assert client.put('/api/assessment/progress',json={'answers':{'1':5}}).status_code==200
    answers={str(i):(i%5)+1 for i in range(1,31)}
    task=client.post('/api/assessment/score',json={'answers':answers,'job_zone':4}).json()
    result=wait(client,task['task_id']);assert result['status']=='completed',result
    assert len(next(p['answers'] for name,p in calls if name=='results'))==30
    assert next(p for name,p in calls if name=='careers')['zone']==4
    assert client.get('/api/workspace').json()['assessment_progress'][0]['data']['answers']==answers

def test_dates_duplicates_and_country():
    current=datetime(2026,10,20,tzinfo=timezone.utc)
    a=integrations.normalize_job('ashby','Junior Data Analyst','Example','Phoenix, AZ','Use SQL.','https://example.com/jobs/1?utm_source=demo',published_at=(current-timedelta(days=2)).isoformat(),publication_source='Ashby last publishedAt')
    assert integrations.recent(a,7,current)
    unknown={**a,'published_at':None,'updated_at':current.isoformat()}
    assert not integrations.recent(unknown,7,current)
    future={**a,'published_at':(current+timedelta(days=1)).isoformat()}
    assert not integrations.recent(future,7,current)
    b=integrations.normalize_job('web','Junior Data Analyst','Example','Phoenix, AZ','Use SQL and Excel.','https://example.com/jobs/1')
    merged=integrations.deduplicate([a,b]);assert len(merged)==1 and len(merged[0]['sources'])==2
    assert integrations.in_us(a)
    assert not integrations.in_us({**a,'country':'Canada'})

def test_search_partial_failures_cache_and_discovery_slot(client,account,monkeypatch):
    calls=[]
    async def board(board,title):
        calls.append(board)
        if board['provider']=='greenhouse':raise ServiceError('Source fixture outage.')
        return [integrations.normalize_job('ashby','Junior Data Analyst','Example','Phoenix, AZ','Use SQL.','https://example.com/jobs/2',country='US',published_at=db.now(),publication_source='Ashby last publishedAt')]
    monkeypatch.setattr(integrations,'board_jobs',board)
    body={'title':'data analyst','sources':['greenhouse','ashby','jobspy']}
    result=wait(client,client.post('/api/jobs/search',json=body).json()['task_id'])
    assert result['status']=='completed',result
    data=result['result']['data'];assert len(data['jobs'])==1
    assert any(s['status']=='error' for s in data['sources'])
    assert any(s['status']=='disabled' for s in data['sources'])
    count=len(calls)
    assert client.post('/api/jobs/search',json=body).json()['cached']
    assert len(calls)==count
    stamp=db.now()
    with db.connect() as c:c.execute('INSERT INTO tasks VALUES(?,?,?,?,?,?,?,?)',('active-fixture',account['id'],'discovery','running',None,None,stamp,stamp))
    assert client.post('/api/jobs/search',json={**body,'refresh':True}).status_code==409

@pytest.mark.parametrize('url',['http://127.0.0.1/a','http://localhost/a','https://user:password@example.com','file:///etc/passwd','http://169.254.169.254/','http://[::1]/','https://example.com:8000/'])
def test_private_urls_rejected(url,monkeypatch):
    monkeypatch.setattr(socket,'getaddrinfo',lambda *a,**k:[(socket.AF_INET,socket.SOCK_STREAM,6,'',('127.0.0.1',80))])
    with pytest.raises(ServiceError):public_address(url)

def test_html_job_structured_date_not_update():
    result=html_job(b'<script type="application/ld+json">{"@type":"JobPosting","title":"Analyst","description":"Use SQL","datePosted":"2026-10-19","hiringOrganization":{"name":"Example"}}</script>')
    assert result['published_at']=='2026-10-19'
    assert result['company']=='Example'

def test_backup_import_and_account_deletion(client,sample,account):
    upload=client.post('/api/documents',files={'file':('source.txt',b'Fictional student source document.','text/plain')}).json()
    raw=client.get('/api/account/export').content
    assert b'fictional-test-password' not in raw and b'csrf' not in raw and b'token_hash' not in raw
    payload=json.loads(raw)
    count=len(payload['records'])
    imported=client.post('/api/account/import',files={'file':('backup.json',raw,'application/json')})
    assert imported.status_code==200,imported.text
    assert imported.json()['imported']==count
    records=client.get('/api/workspace').json()
    assert len(records['experience'])==6
    bad={**payload,'schema_version':99}
    assert client.post('/api/account/import',files={'file':('bad.json',json.dumps(bad).encode())}).status_code==400
    assert client.request('DELETE','/api/account',json={'username':'student','password':'wrong'}).status_code==401
    assert client.request('DELETE','/api/account',json={'username':'student','password':'fictional-test-password'}).status_code==200
    assert client.get('/api/me').status_code==401
    assert not (config.DATA_DIR/'uploads'/account['id']).exists()
    with db.connect() as c:
        assert c.execute('SELECT count(*) FROM entities').fetchone()[0]==0
        assert c.execute('SELECT count(*) FROM evidence_fts WHERE user_id=?',(account['id'],)).fetchone()[0]==0


def test_structured_job_geography():
    result=html_job(b'<script type="application/ld+json">{"@type":"JobPosting","jobLocation":{"address":{"addressLocality":"Tucson","addressRegion":"AZ","addressCountry":"US"}},"jobLocationType":"TELECOMMUTE"}</script>')
    assert result['country']=='US'
    assert result['location']=='Tucson, AZ, US'
    assert result['workplace']=='remote'
