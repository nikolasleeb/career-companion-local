import io
import os
import re
import shutil
import subprocess
import tempfile
import threading
from functools import lru_cache
from pathlib import Path
from zipfile import ZipFile, BadZipFile
from xml.sax.saxutils import escape
from defusedxml import ElementTree
from pypdf import PdfReader
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from . import config, resume_presentation
from .network import ServiceError

_converter_slots=threading.BoundedSemaphore(2)

def docx_pdf(raw):
    if not _converter_slots.acquire(timeout=30):
        raise ServiceError('Document conversion is busy. Wait for a current conversion to finish and retry.')
    try:return _docx_pdf(raw)
    finally:_converter_slots.release()

def _docx_pdf(raw):
    executable = os.getenv('SOFFICE_BIN') or shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        raise ServiceError('DOCX page validation requires the document converter included in the Docker image. Use the Docker app or paste plain text.')
    with tempfile.TemporaryDirectory(prefix='nextstep-doc-') as folder:
        root = Path(folder)
        source = root/'input.docx'
        source.write_bytes(raw)
        profile = root/'profile'
        registry = profile/'user'
        registry.mkdir(parents=True)
        (registry/'registrymodifications.xcu').write_text('''<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry"><item oor:path="/org.openoffice.Office.Common/Security/Scripting"><prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item></oor:items>''')
        try:
            subprocess.run([executable, '-env:UserInstallation='+profile.as_uri(), '--headless','--nologo','--nodefault','--norestore','--convert-to','pdf','--outdir',str(root),str(source)], capture_output=True, timeout=45, check=True, env={'PATH':os.defpath,'LANG':'C.UTF-8','HOME':str(root),'XDG_CACHE_HOME':str(root/'cache'),'TMPDIR':str(root)})
            return (root/'input.pdf').read_bytes()
        except (subprocess.SubprocessError,OSError):
            raise ServiceError('The DOCX converter could not safely finish within 45 seconds. Paste the text or use a text-based PDF.') from None

def docx_pages(raw):
    return len(PdfReader(io.BytesIO(docx_pdf(raw))).pages)

def enrich_docx_links(text,raw):
    try:
        doc=Document(io.BytesIO(raw))
        for paragraph in doc.paragraphs:
            for link in paragraph.hyperlinks:
                if link.url.startswith(('http://','https://')) and link.text and link.text in text and link.url not in text:
                    text=text.replace(link.text,link.text+' ('+link.url+')',1)
    except Exception:
        pass  # The original extracted text is still available for review.
    return text

def extract(filename, raw, *, max_upload=None, repair_font_mappings=False):
    limit=config.MAX_UPLOAD if max_upload is None else max_upload
    if not raw or len(raw)>limit:
        raise ServiceError(f'Choose a nonempty document of {limit//(1024*1024)} MB or less.')
    suffix = Path(filename).suffix.lower()
    try:
        if suffix=='.pdf':
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted:
                raise ServiceError('This PDF is encrypted. Upload an unlocked document or paste text.')
            pages = len(reader.pages)
            if pages>30:
                raise ServiceError('The document exceeds the 30-page limit. Upload the relevant pages.')
            if repair_font_mappings:
                from .pdf_text import repair_ligatures
                repair_ligatures(reader)
            text = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
            if len(text.strip())<30:
                raise ServiceError('This PDF appears to be scanned or contains too little selectable text. Paste the text or upload a text-based PDF.')
        elif suffix=='.docx':
            with ZipFile(io.BytesIO(raw)) as archive:
                if len(archive.infolist())>2000 or sum(i.file_size for i in archive.infolist())>40*1024*1024:
                    raise ServiceError('This DOCX expands beyond the safe document size limit.')
                if any('vbaProject' in i.filename or i.flag_bits & 1 for i in archive.infolist()):
                    raise ServiceError('Macro-enabled or encrypted documents are not supported.')
                names=[item.filename for item in archive.infolist()]
                if len(set(names))!=len(names) or any(name.startswith(('/', '\\')) or '..' in name.replace('\\','/').split('/') for name in names):
                    raise ServiceError('This DOCX contains unsafe archive paths or duplicate entries.')
                for name in names:
                    if not name.endswith('.rels'):continue
                    relationships=ElementTree.fromstring(archive.read(name))
                    for relationship in relationships:
                        if relationship.get('TargetMode')=='External' and not relationship.get('Type','').endswith('/hyperlink'):
                            raise ServiceError('Externally linked images, templates and embedded resources are not supported. Export a self-contained document.')
                xml = ElementTree.fromstring(archive.read('word/document.xml'))
                # Reject external embedded objects; ordinary hyperlinks remain text.
                if any(e.tag.endswith('}altChunk') or e.tag.endswith('}object') for e in xml.iter()):
                    raise ServiceError('Embedded objects are not supported. Export a simple document or paste text.')
            pages = docx_pages(raw)
            if pages>30:
                raise ServiceError('The document exceeds the 30-page limit. Upload the relevant pages.')
            doc = Document(io.BytesIO(raw))
            def paragraph_text(paragraph):
                value=paragraph.text
                # Preserve actual DOCX hyperlink destinations instead of just
                # labels such as LinkedIn. Links are never fetched.
                for link in paragraph.hyperlinks:
                    if link.url.startswith(('http://','https://')) and link.text and link.url not in value:
                        value=value.replace(link.text,link.text+' ('+link.url+')',1)
                props=paragraph._p.pPr
                numbered=props is not None and props.numPr is not None
                if numbered or paragraph.style.name.lower().startswith('list bullet'):
                    value='• '+value
                return value
            text = '\n'.join(paragraph_text(p) for p in doc.paragraphs)
            text += '\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
            for section in doc.sections:
                text += '\n'+'\n'.join(p.text for part in (section.header,section.footer) for p in part.paragraphs)
        elif suffix=='.txt':
            pages = None
            text = raw.decode('utf-8-sig')
            if '\x00' in text:
                raise ServiceError('This file is not readable plain text.')
        else:
            raise ServiceError('Supported document formats are PDF, DOCX, and UTF-8 TXT.')
    except ServiceError:
        raise
    except (Exception,):
        raise ServiceError('This document could not be read. Try an unlocked PDF, a simple DOCX, or pasted text.') from None
    if len(text)>config.MAX_TEXT:
        raise ServiceError('The extracted text exceeds 100,000 characters. Upload only the relevant sections.')
    if not text.strip():
        raise ServiceError('No readable text was found in this document.')
    return {'text':text.strip(),'pages':pages,'format':suffix[1:]}

def clean(value):
    value = str(value or '').strip()
    if re.search(r'\[(?:missing|todo|insert|date|name|company)[^\]]*\]|\{\{.*?\}\}',value,re.I):
        return ''
    return value

def resume_plain(data):
    data=resume_presentation.display(data)
    lines = [clean(data.get('name')) or 'Resume']
    if clean(data.get('headline')): lines.append(clean(data['headline']))
    lines.extend(contact.strip() for contact in data.get('contact_lines',str(data.get('contact','')).splitlines()) if contact.strip())
    for section in data['sections']:
        if not section['entries']: continue
        lines += ['',section['title'].upper()]
        for entry in section['entries']:
            lines += [' | '.join(filter(None,[clean(entry.get('title')),clean(entry.get('details')) or ' | '.join(filter(None,[clean(entry.get('organization')),clean(entry.get('dates'))]))]))]
            lines += ['- '+clean(b) for b in entry.get('bullets',[]) if clean(b)]
            lines.append('')
    return '\n'.join(lines).strip()+'\n'

def export_resume(data, format):
    if data.get('artifacts',{}).get(format):
        import base64
        return base64.b64decode(data['artifacts'][format]), {'pdf':'application/pdf','docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}[format]
    if format=='txt':
        return resume_plain(data).encode(), 'text/plain; charset=utf-8'
    data=resume_presentation.display(data)
    buffer = io.BytesIO()
    if format=='docx':
        doc = Document()
        section = doc.sections[0]
        section.page_width,section.page_height = Inches(8.5), Inches(11)
        section.top_margin=section.bottom_margin=Inches(.65 if data.get('schema_version',0)>=2 else .7)
        section.left_margin=section.right_margin=Inches(.7 if data.get('schema_version',0)>=4 else .8)
        for key in ('Normal','Title','Heading 1','Heading 2','List Bullet'):
            style=doc.styles[key]; style.font.name='Arial'; style.font.color.rgb=RGBColor(0,0,0)
            if data.get('schema_version',0)>=2:
                style.element.get_or_add_rPr().append(OxmlElement('w:kern'))
                style.element.rPr[-1].set(qn('w:val'),'0')
            style.font.size=Pt(10.5 if key in ('Normal','List Bullet') else 12 if key.startswith('Heading') else 21)
            style.paragraph_format.space_after=Pt(2 if data.get('schema_version',0)>=2 else 5)
            style.paragraph_format.line_spacing=1.03 if data.get('schema_version',0)>=4 else 1 if data.get('schema_version',0)>=2 else 1.1
        # Override theme colors and the built-in title paragraph border.
        title=doc.add_paragraph(clean(data.get('name')) or 'Resume','Title')
        title.paragraph_format.space_after=Pt(6)
        for ppr in (title._p.get_or_add_pPr(),doc.styles['Title'].element.get_or_add_pPr()):
            for border in list(ppr.findall(qn('w:pBdr'))): ppr.remove(border)
        if clean(data.get('headline')):
            headline=doc.add_paragraph(clean(data['headline']))
            if data.get('schema_version',0)>=4:
                headline.paragraph_format.space_after=Pt(4)
                for run in headline.runs:run.bold=True;run.font.size=Pt(11)
        if data.get('schema_version',0)>=4:
            bullets=doc.styles['List Bullet'].paragraph_format
            bullets.left_indent=Inches(.18);bullets.first_line_indent=Inches(-.13)
            bullets.space_after=Pt(3)
        for contact in data.get('contact_lines',str(data.get('contact','')).splitlines()):
            if clean(contact): doc.add_paragraph(clean(contact))
        for block in data['sections']:
            if not block['entries']: continue
            heading=doc.add_paragraph(block['title'],'Heading 1'); heading.paragraph_format.space_before=Pt(7 if data.get('schema_version',0)>=2 else 13)
            for entry in block['entries']:
                modern=data.get('schema_version',0)>=4
                p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
                p.add_run(clean(entry['title'])).bold=True
                if modern:
                    p.paragraph_format.space_after=Pt(1)
                    employer=' · '.join(filter(None,[clean(entry.get('organization')),clean(entry.get('location'))]))
                    arrangement={'onsite':'On-site','hybrid':'Hybrid','remote':'Remote'}.get(entry.get('work_arrangement'),'')
                    if arrangement:employer+=' · '+arrangement if employer else arrangement
                    dates=clean(entry.get('dates'))
                    if employer or dates:
                        metadata=doc.add_paragraph();metadata.paragraph_format.space_after=Pt(2)
                        metadata.paragraph_format.keep_with_next=bool(entry.get('bullets'))
                        if employer:metadata.add_run(employer)
                        if dates:
                            if len(employer)+len(dates)>102:
                                if employer:metadata.add_run('\n')
                                metadata.add_run(dates)
                            elif employer:
                                metadata.paragraph_format.tab_stops.add_tab_stop(Inches(6.9),WD_TAB_ALIGNMENT.RIGHT)
                                metadata.add_run('\t'+dates)
                            else:metadata.add_run(dates)
                else:
                    details=clean(entry.get('details')) or ' | '.join(filter(None,[clean(entry.get('organization')),clean(entry.get('dates'))]))
                    if details:p.add_run(' | '+details)
                if data.get('schema_version',0)>=3 and entry.get('category')=='skills':
                    if entry.get('bullets'):p.add_run(': '+', '.join(clean(v) for v in entry['bullets']))
                    p.paragraph_format.keep_with_next=False
                else:
                    courses=[clean(b) for b in entry.get('bullets',[]) if b.startswith('Relevant class: ')] if data.get('schema_version',0)>=3 else []
                    for bullet in entry.get('bullets',[]):
                        if clean(bullet) and bullet not in courses:
                            paragraph=doc.add_paragraph(clean(bullet),'List Bullet')
                            if modern:paragraph.paragraph_format.keep_together=True
                    if courses:
                        paragraph=doc.add_paragraph('Relevant courses: '+', '.join(v.removeprefix('Relevant class: ') for v in courses),'List Bullet')
                        if modern:paragraph.paragraph_format.keep_together=True
        if data.get('schema_version',0)>=2:
            for paragraph in doc.paragraphs:
                for run in paragraph.runs:
                    props=run._r.get_or_add_rPr()
                    for existing in list(props.findall(qn('w:kern'))): props.remove(existing)
                    kern=OxmlElement('w:kern');kern.set(qn('w:val'),'0');props.append(kern)
        doc.core_properties.author=''
        doc.core_properties.title='Resume'
        doc.save(buffer)
        return buffer.getvalue(),'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    if format!='pdf': raise ServiceError('Choose PDF, DOCX, or TXT.')
    font='Helvetica'
    bold='Helvetica-Bold'
    candidates=[Path('/usr/share/fonts/truetype/dejavu'),Path(os.getenv('FONT_DIR',str(Path.home()/'.local/share/fonts')))]
    for path in candidates:
        normal=path/'DejaVuSans.ttf'; weight=path/'DejaVuSans-Bold.ttf'
        if normal.exists() and weight.exists():
            if 'ResumeSans' not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont('ResumeSans',str(normal))); pdfmetrics.registerFont(TTFont('ResumeSansBold',str(weight)))
            font,bold='ResumeSans','ResumeSansBold'
            break
    normal=ParagraphStyle('body',fontName=font,fontSize=10.5,leading=14,spaceAfter=4,splitLongWords=True)
    title_style=ParagraphStyle('title',parent=normal,fontName=bold,fontSize=21,leading=27,spaceAfter=7)
    heading_style=ParagraphStyle('heading',parent=normal,fontName=bold,fontSize=12,leading=17,spaceBefore=10,spaceAfter=5,keepWithNext=True)
    entry_style=ParagraphStyle('entry',parent=normal,fontName=bold,keepWithNext=True)
    bullet_style=ParagraphStyle('bullet',parent=normal,leftIndent=12,firstLineIndent=-10,spaceAfter=4)
    story=[Paragraph(escape(clean(data.get('name')) or 'Resume'),title_style)]
    if clean(data.get('headline')): story.append(Paragraph(escape(clean(data['headline'])),normal))
    for contact in data.get('contact_lines',str(data.get('contact','')).splitlines()):
        if clean(contact): story.append(Paragraph(escape(clean(contact)),normal))
    for block in data['sections']:
        if not block['entries']: continue
        story.append(Paragraph(escape(block['title']),heading_style))
        for entry in block['entries']:
            text=' | '.join(filter(None,[clean(entry['title']),clean(entry.get('details')) or ' | '.join(filter(None,[clean(entry.get('organization')),clean(entry.get('dates'))]))]))
            story.append(Paragraph(escape(text),entry_style))
            for bullet in entry.get('bullets',[]):
                if clean(bullet): story.append(Paragraph('• '+escape(clean(bullet)).replace('\n','<br/>'),bullet_style))
            story.append(Spacer(1,3))
    SimpleDocTemplate(buffer,pagesize=letter,leftMargin=57.6,rightMargin=57.6,topMargin=50.4,bottomMargin=50.4,title='Resume',author='').build(story)
    return buffer.getvalue(),'application/pdf'

@lru_cache(maxsize=32)
def pdf_preview(raw,page=1):
    executable=shutil.which('pdftoppm')
    if not executable: raise ServiceError('PDF preview needs the renderer included in the Docker app. Downloads remain available.')
    with tempfile.TemporaryDirectory(prefix='nextstep-preview-') as folder:
        source=Path(folder)/'resume.pdf';source.write_bytes(raw)
        target=Path(folder)/'page'
        try:
            subprocess.run([executable,'-f',str(page),'-singlefile','-scale-to','1500','-png',str(source),str(target)],capture_output=True,check=True,timeout=30)
            return target.with_suffix('.png').read_bytes()
        except (OSError,subprocess.SubprocessError): raise ServiceError('The saved PDF could not be rendered for preview. Download it to inspect the layout.') from None
